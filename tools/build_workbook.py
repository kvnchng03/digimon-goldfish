"""Build the Glowing Dawn theorycraft workbook.

    python3 tools/build_workbook.py [--out reports/glowing-dawn-theorycraft.xlsx] [--games 20000]

Simulation numbers are produced live by the goldfish engine while the script runs;
the consistency sheet is pure Excel formulas over the decklist; the matchup sheet
computes per-check probabilities from the opponents' 50-card lists (regional top
cut, 2026-10-04). Research-derived data (timings, win rates, play rules, playbooks,
each tagged OBSERVED / STATED / INFERRED / COUNTED / SIM) lives in research_data.py.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import random
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from research_data import (  # noqa: E402
    EVENT_RESULTS,
    FLAWS,
    PLAYBOOKS,
    PROFILE_SECTIONS,
    RESEARCH_SOURCES,
    RULES,
    SECURITY_LINKS,
    TAGS,
    VIDEOS,
    WINRATES,
)
from xlsx_style import (  # noqa: E402
    F_BASE,
    F_BOLD,
    F_INPUT,
    F_LINK,
    F_NOTE,
    F_TAG,
    FILL_KEY,
    FILL_SUB,
    NUM1,
    NUM2,
    PCT,
    finalize,
    header,
    link,
    note,
    put,
    title,
    widths,
)

from goldfish import CARDS, Config, apply_edits, parse_decklist  # noqa: E402
from goldfish import calc as gc  # noqa: E402
from goldfish.sim import run, summarize  # noqa: E402


def put_tag(ws, row, col, tag, source=None):
    """An evidence-tag cell, filled by its first tag word, with the source as a comment."""
    fill = TAGS.get(tag.split()[0].rstrip(","), (None, None))[1]
    return put(ws, row, col, tag, F_TAG, fill=PatternFill("solid", fgColor=fill) if fill else None, wrap=True,
               comment=source or None)



# ---------------------------------------------------------------------- sheets
GUIDE_ROUTINE = [
    ("Before an event", "Read the four matchup sheets and Play Rules. Check your list against the 4th-place list.", "vs Jupitermon"),
    ("Before each round", "Reread the sheet for your next opponent. Know their big turn (their turn 2), your Lv6, and how much memory to pass them.", "vs Toho"),
    ("During a game", "Not sure whether to attack? Enter the board in the Calculator and look at the shaded 'Plan for' chance.", "Calculator"),
    ("After a game", "Replay the turn that decided it in the Calculator. Was there a line with a better 'Plan for' chance? That's the lesson.", "Calculator"),
    ("Changing the deck", "Change the blue counts on Decklist and read Consistency. Ask for a simulator rerun to update Goldfish.", "Decklist"),
]
GUIDE_DRILLS = [
    ("1. Opening hands", "Shuffle, draw 5, decide: keep or mulligan? Keep a Lv3 plus a Tamer or a way to find one (e-Pulse, Gekkomon, Liollmon). Do 20 hands, then compare how often you saw a Tamer or e-Pulse with the Consistency sheet."),
    ("2. Goldfish yourself", "Play turns 1-4 alone with no opponent and count the turn you could deal 6 hits. The simulator wins by turn 4 about 77% of the time. If you're slower, compare your choices with Play Rules."),
    ("3. Guess, then check", "Set up a board, guess your chance to win this turn, then check the Calculator. Your guesses get sharper fast, and that's the skill you use at the table."),
    ("4. Know their turn 2", "For each meta deck, say what it does on its turn 2 and what you do about it. Check yourself against the top box of its sheet."),
    ("5. Exact memory", "Practice ending turns on an exact pass: 4 against Homeros, as low as you can against the others, with 1-2 spare on your kill turn."),
    ("6. Review losses", "Write down the turn you lost and why. Test the other line in the Calculator, and check whether a Play Rule covered it."),
]
GUIDE_RULES = [
    "Land a Tamer on turn 1-2. Mulligan hands that can't find one.",
    "Don't chip. Go for the kill or hold back: a poke gives them free Tamers 20% of the time.",
    "Leave 1-2 spare memory on your kill turn in case a Tamer flips from their security.",
    "Atratusmon is your Lv6 in every matchup. Attack with it first, so it's immune before their triggers fire.",
    "Pass Homeros exactly 4 memory. End your turn with Jupitermon at 0 security.",
    "Against Toho and DATS, keep your Lv4 in the breeding area on turn 2: both delete it on sight.",
]
GUIDE_SHEETS = [
    ("Calculator", "Your chance to win this turn, given their security, blockers and your attackers."),
    ("vs Jupitermon", "Matchup guide: their plan, how they beat you, your plan turn by turn, their mistakes to punish."),
    ("vs Homeros", "Same, for TS Mervamon (Homeros control)."),
    ("vs Toho", "Same, for Toho Braves."),
    ("vs DATS", "Same, for DATA SQUAD Ravemon (Rose Rave)."),
    ("Play Rules", "Do's and don'ts, each with the number or game behind it, and the deck's known weak spots."),
    ("Decklist", "Your 50 + 4. Change the blue counts to try a different list."),
    ("Consistency", "Odds of drawing what you need (a Tamer, a Lv3, a Lv6) in a given number of cards."),
    ("Goldfish", "How fast your deck kills with no opponent, from 20,000 simulated games."),
    ("Board Setup", "What your board usually looks like at the end of each turn."),
    ("Fuel and Memory", "How much fuel and memory each card makes, and the main combo step by step."),
    ("Deck Comparison", "Your list next to the 4th-place Dusseldorf list, and why each difference matters."),
    ("Matchups", "Real win rates with their likely range, and what each deck's security does to your attackers."),
    ("Opponent Profiles", "What each meta deck does on each of its turns. The practice opponents will be built from this."),
    ("Meta Decklists", "The regional top-cut lists everything is based on."),
    ("Sources", "Where every number came from, including the videos."),
]


def sheet_readme(wb, games, seed, stamp):
    ws = wb.active
    ws.title = "Start Here"
    title(ws, "Start here", f"How to read this workbook and use it to get better. Built {stamp}.")
    r = 4
    header(ws, r, 1, ["Your routine", "What to do", "Open"])
    for when, what, sheet in GUIDE_ROUTINE:
        r += 1
        put(ws, r, 1, when, F_BOLD)
        put(ws, r, 2, what, wrap=True)
        link(ws, r, 3, sheet)
    r += 2
    header(ws, r, 1, ["Practice drills", "How", ""])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    for name, how in GUIDE_DRILLS:
        r += 1
        put(ws, r, 1, name, F_BOLD)
        put(ws, r, 2, how, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 2
    header(ws, r, 1, ["Six rules to remember", "", ""])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    link(ws, r, 3, "Play Rules", "All rules")
    for i, rule in enumerate(GUIDE_RULES, 1):
        r += 1
        put(ws, r, 1, i, F_BOLD, align="center")
        put(ws, r, 2, rule, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 2
    header(ws, r, 1, ["Every sheet (click to open)", "What it's for", ""])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    for sheet, what in GUIDE_SHEETS:
        r += 1
        link(ws, r, 1, sheet)
        put(ws, r, 2, what, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 2
    header(ws, r, 1, ["Colors", "Meaning", ""])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    legend = [("Blue text", F_INPUT, None, "A value you can change."),
              ("Green text", F_LINK, None, "Pulled from another sheet."),
              ("Shaded", F_BASE, FILL_KEY, "The number that matters most.")]
    for text, font, fill, meaning in legend:
        r += 1
        put(ws, r, 1, text, font, fill=fill)
        put(ws, r, 2, meaning)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    for tag, (meaning, _) in TAGS.items():
        r += 1
        put_tag(ws, r, 1, tag)
        put(ws, r, 2, meaning)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 2
    header(ws, r, 1, ["Good to know", "", ""])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for text in (
        "Turn numbers are each player's own turns: 'their T2' is the opponent's second turn.",
        "Goldfish numbers are your deck alone, with no opponent. They show speed and consistency, not win rate.",
        "Hover over an evidence tag to see where it came from.",
        "Still to come: practice opponents for each meta deck, and a fuel check in the Calculator.",
    ):
        r += 1
        put(ws, r, 1, text, wrap=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    widths(ws, {"A": 26, "B": 92, "C": 18})


ROLE_NOTES = {
    "ST23-01": "Egg. Inherited: when attacking, 1 fuel -> digivolve from hand at -2. The engine's cheapest digivolve.",
    "ST23-06": "On play / when moving: reveal 3, add 1, fuel 1. Inherited Piercing. Best Lv3 to move out.",
    "BT25-032": "On play: reveal 3, add 1 + 1 yellow. Inherited Barrier (security-for-life in battle).",
    "BT26-025": "On play / when moving: top security -> fuel, Recovery +1. Inherited: security to hand when attacking.",
    "ST23-12": "On play: 1 fuel -> Glowing Dawn Digimon from trash to hand. Inherited Retaliation.",
    "ST23-03": "Security -> hand, Recovery +1; its own digivolve costs -2 for 1 fuel. Inherited Barrier. Best Lv4 base.",
    "BT25-035": "-3000 DP; 2 fuel -> free digivolve into a Lv5 from hand. Inherited Barrier.",
    "BT26-026": "Barrier; when attacking, 1 fuel or 1 security -> Glowing Dawn Option at -2. Inherited Barrier.",
    "BT25-049": "Suspend 1; Glowing Dawn Options cost -3 for 1 fuel (e-Pulse becomes free). Inherited Piercing.",
    "BT26-070": "Draw 1 trash 1; [Main] 2 fuel -> reuse a Glowing Dawn Option from trash at -2.",
    "ST23-04": "Alliance; -5000 DP; 1 fuel -> play/use a Glowing Dawn card at -3. Inherited end-of-attack unsuspend.",
    "BT25-041": "Alliance; when digivolving OR attacking: security->hand or 1 fuel -> play/use at -3. Inherited unsuspend.",
    "BT26-031": "DUAL. Option half -8000 DP (-13000 with a security). No inherited effect: a Lv6 on top gets no unsuspend.",
    "ST23-08": "Alliance; +3000 DP; 1 fuel -> play/use at -3. Inherited end-of-attack unsuspend.",
    "BT25-057": "DUAL. Option half: Rush, Security Attack +1, +5000 DP (your turn only, errata). No inherited effect.",
    "BT25-043": "DUAL Lv6. Recovery +1 then trash top security of whoever has most -> unsuspend. Protects Glowing Dawn Digimon from leaving.",
    "ST23-09": "DUAL Lv6. Security Attack +1, Reboot, Blocker; immune to their Digimon effects until their turn ends; delete lowest DP.",
    "ST23-13": "Start of main / on play: +1 fuel, +1 memory once they have a Digimon. The core Tamer.",
    "BT25-090": "Memory floor 3; +2 fuel whenever any Digimon suspends; Glowing Dawn Options -1.",
    "BT26-089": "Tuck 1 -> draw 1, +1 memory; +1 fuel whenever your security is removed from.",
    "P-236": "Reveal 3, add 1; Delay next turn: +2 memory. Not released in EN as of the wiki snapshot.",
    "ST23-15": "Play a <=4-cost BEATBREAK card from hand or trash for free; next turn tucks for draw 1 +1 memory.",
}


def sheet_decklist(wb, deck):
    ws = wb.create_sheet("Decklist")
    title(ws, "Your list", "Change the blue counts and the Consistency sheet updates.")
    cols = ["Count", "Card", "ID", "Kind", "Level", "Colors", "Play cost", "Option cost", "Digivolve cost (trait route)", "DP", "What it does for the engine"]
    header(ws, 4, 1, cols)
    counts = deck.counts()
    order = sorted(counts, key=lambda cid: ((CARDS[cid].level or 9), {"egg": 0, "digimon": 1, "dual": 2, "tamer": 3, "option": 4}[CARDS[cid].kind.value], cid))
    r = 5
    first = r
    for cid in order:
        c = CARDS[cid]
        put(ws, r, 1, counts[cid], F_INPUT)
        put(ws, r, 2, c.name)
        put(ws, r, 3, cid)
        put(ws, r, 4, c.kind.value)
        put(ws, r, 5, c.level)
        put(ws, r, 6, "/".join(c.colors))
        put(ws, r, 7, c.play_cost)
        put(ws, r, 8, c.option_cost)
        put(ws, r, 9, c.evo_cost)
        put(ws, r, 10, c.dp or None)
        put(ws, r, 11, ROLE_NOTES.get(cid, ""), wrap=True)
        r += 1
    last = r - 1
    rng = lambda col: f"${col}${first}:${col}${last}"  # noqa: E731
    r += 1
    header(ws, r, 2, ["Totals", "Value", "Note"], fill=FILL_SUB, font=F_BOLD)
    r += 1
    totals = [
        ("Main deck cards", f'=SUMIF({rng("D")},"<>egg",{rng("A")})', "must be 50"),
        ("Digi-Eggs", f'=SUMIF({rng("D")},"egg",{rng("A")})', "max 5"),
        ("Lv3", f"=SUMIF({rng('E')},3,{rng('A')})", ""),
        ("Lv4", f"=SUMIF({rng('E')},4,{rng('A')})", ""),
        ("Lv5", f"=SUMIF({rng('E')},5,{rng('A')})", ""),
        ("Lv5 non-DUAL (carry the inherited unsuspend)", f'=SUMIFS({rng("A")},{rng("E")},5,{rng("D")},"digimon")', ""),
        ("Lv5 DUAL", f'=SUMIFS({rng("A")},{rng("E")},5,{rng("D")},"dual")', ""),
        ("Lv6", f"=SUMIF({rng('E')},6,{rng('A')})", ""),
        ("Tamers", f'=SUMIF({rng("D")},"tamer",{rng("A")})', ""),
        ("Tomoro & Kyo", f'=SUMIF({rng("C")},"ST23-13",{rng("A")})', ""),
        ("Options", f'=SUMIF({rng("D")},"option",{rng("A")})', ""),
        ("e-Pulse", f'=SUMIF({rng("C")},"ST23-15",{rng("A")})', ""),
        ("Engine pieces (Tamers + e-Pulse)", None, "sum of the two rows above"),
        ("Barrier bases (Liollmon BT25, all Cougarmon)", f'=SUMIF({rng("C")},"BT25-032",{rng("A")})+SUMIF({rng("C")},"ST23-03",{rng("A")})+SUMIF({rng("C")},"BT25-035",{rng("A")})+SUMIF({rng("C")},"BT26-026",{rng("A")})', "stacks that can trade a security card for the Digimon's life"),
    ]
    total_rows = {}
    for label, formula, n in totals:
        put(ws, r, 2, label, F_BOLD, wrap=True)
        if formula is None:
            formula = f"=C{total_rows['Tamers']}+C{total_rows['e-Pulse']}"
        put(ws, r, 3, formula)
        note(ws, r, 4, n)
        total_rows[label] = r
        r += 1
    widths(ws, {"A": 10, "B": 30, "C": 11, "D": 9, "E": 7, "F": 13, "G": 10, "H": 11, "I": 14, "J": 8, "K": 95})
    return total_rows


def sheet_consistency(wb, total_rows):
    ws = wb.create_sheet("Consistency")
    title(ws, "Draw odds", "Chance to see at least one copy among the cards you've seen. Counts come from the Decklist sheet.")
    put(ws, 4, 1, "Deck size N", F_BOLD)
    put(ws, 4, 2, f"=Decklist!$C${total_rows['Main deck cards']}", F_LINK)
    put(ws, 5, 1, "Library after setup (50 - 5 hand - 5 security)", F_BOLD)
    put(ws, 5, 2, "=B4-10")
    seen = [5, 6, 7, 8, 10, 12]
    header(ws, 7, 1, ["Event", "Copies K"] + [f"seen {n}" for n in seen] + ["What 'seen' means"])
    events = [
        ("Any Lv3", f"=Decklist!$C${total_rows['Lv3']}", "5 = opening hand; 6 = turn 1 going second"),
        ("Any Lv4", f"=Decklist!$C${total_rows['Lv4']}", "7 to 8 = by turn 2 incl. the digivolution draw"),
        ("Any Lv5", f"=Decklist!$C${total_rows['Lv5']}", "10 = by turn 3"),
        ("Any Lv6", f"=Decklist!$C${total_rows['Lv6']}", "10 = by turn 4, 12 = by turn 5"),
        ("Any Tamer", f"=Decklist!$C${total_rows['Tamers']}", ""),
        ("Tamer or e-Pulse", f"=Decklist!$C${total_rows['Engine pieces (Tamers + e-Pulse)']}", "an engine piece"),
        ("Tomoro & Kyo specifically", f"=Decklist!$C${total_rows['Tomoro & Kyo']}", ""),
    ]
    r = 8
    ev_rows = {}
    for label, k, what in events:
        put(ws, r, 1, label, F_BOLD)
        put(ws, r, 2, k, F_LINK)
        for j, n in enumerate(seen):
            put(ws, r, 3 + j, f"=1-COMBIN($B$4-$B{r},{n})/COMBIN($B$4,{n})", fmt=PCT)
        note(ws, r, 3 + len(seen), what)
        ev_rows[label] = r
        r += 1
    r += 1
    header(ws, r, 1, ["Derived", "Value", "Formula"], fill=FILL_SUB, font=F_BOLD)
    r += 1
    lv3 = ev_rows["Any Lv3"]
    derived = [
        ("Lv3 in opening 5, mulligan if none", f"=1-(1-C{lv3})^2", "1 - (1 - p5)^2", True),
        ("Lv3 AND an engine piece in opening 5",
         f"=1-(COMBIN($B$4-$B{lv3},5)+COMBIN($B$4-$B{ev_rows['Tamer or e-Pulse']},5)-COMBIN($B$4-$B{lv3}-$B{ev_rows['Tamer or e-Pulse']},5))/COMBIN($B$4,5)",
         "inclusion-exclusion", True),
        ("Lv3 AND an engine piece in 6 (going second)",
         f"=1-(COMBIN($B$4-$B{lv3},6)+COMBIN($B$4-$B{ev_rows['Tamer or e-Pulse']},6)-COMBIN($B$4-$B{lv3}-$B{ev_rows['Tamer or e-Pulse']},6))/COMBIN($B$4,6)",
         "", True),
        ("2 or more Lv6 in the opening 5 (dead cards)",
         f"=1-COMBIN($B$4-$B{ev_rows['Any Lv6']},5)/COMBIN($B$4,5)-$B{ev_rows['Any Lv6']}*COMBIN($B$4-$B{ev_rows['Any Lv6']},4)/COMBIN($B$4,5)",
         "1 - P(0) - P(1)", True),
        ("Reveal 3 from the library finds a Lv4", f"=1-COMBIN($B$5-$B{ev_rows['Any Lv4']},3)/COMBIN($B$5,3)", "Gekkomon / Liollmon / Glowing Dawn", True),
        ("Reveal 3 from the library finds a Lv6", f"=1-COMBIN($B$5-$B{ev_rows['Any Lv6']},3)/COMBIN($B$5,3)", "", True),
        ("Reveal 3 from the library finds a Tamer", f"=1-COMBIN($B$5-$B{ev_rows['Any Tamer']},3)/COMBIN($B$5,3)", "", True),
    ]
    for label, f, n, pct in derived:
        put(ws, r, 1, label, F_BOLD)
        put(ws, r, 2, f, fmt=PCT if pct else None)
        note(ws, r, 3, n)
        r += 1
    r += 1
    note(ws, r, 1, "Approximation: the library-reveal rows assume the remaining copies are all still in the 40-card library.")
    widths(ws, {"A": 46, "B": 11, "C": 10, "D": 10, "E": 10, "F": 10, "G": 10, "H": 10, "I": 50})


def run_sims(deck, games, seed):
    cfg = Config()
    out = {}
    out["random"] = summarize(run(deck, games, seed, cfg, "random"), cfg)
    out["first"] = summarize(run(deck, games, seed, cfg, "first"), cfg)
    out["second"] = summarize(run(deck, games, seed, cfg, "second"), cfg)
    half = max(games // 2, 2000)
    for t in (1, 3, 99):
        c = Config(opp_board_turn=t)
        out[f"opp{t}"] = summarize(run(deck, half, seed, c, "random"), c)
    c5 = Config(max_give=5)
    out["give5"] = summarize(run(deck, half, seed, c5, "random"), c5)
    edited = apply_edits(deck, "BT26-089:+1,ST23-15:-1")
    out["edit_base"] = summarize(run(deck, half, seed, cfg, "random"), cfg)
    out["edit_kyo"] = summarize(run(edited, half, seed, cfg, "random"), cfg)
    out["games"] = games
    out["half"] = half
    return out


def sheet_goldfish(wb, sims, seed):
    ws = wb.create_sheet("Goldfish")
    g, h = sims["games"], sims["half"]
    title(ws, "Goldfish results", f"Your deck played {g:,} times against an opponent who does nothing. Numbers move by about {(0.25 / g) ** 0.5 * 100:.1f} points between runs.")
    turns = [3, 4, 5, 6, 7, 8]
    r = 4
    header(ws, r, 1, ["Kill by end of turn (6 hits)"] + [f"T{t}" for t in turns] + ["mean kill turn", "no kill by T10", "deck-out"])
    for key, label in (("random", "random first/second"), ("first", "going first"), ("second", "going second")):
        r += 1
        s = sims[key]
        put(ws, r, 1, label, F_BOLD)
        for j, t in enumerate(turns):
            put(ws, r, 2 + j, s.kill_by[t], fmt=PCT, fill=FILL_KEY if (key == "random" and t in (4, 5)) else None)
        put(ws, r, 2 + len(turns), s.kill_mean, fmt=NUM2)
        put(ws, r, 3 + len(turns), s.no_kill, fmt=PCT)
        put(ws, r, 4 + len(turns), s.deck_out, fmt=PCT)
    r += 2
    header(ws, r, 1, ["Cumulative checks dealt (mean)"] + [f"T{t}" for t in turns])
    for key, label in (("random", "random"), ("first", "first"), ("second", "second")):
        r += 1
        put(ws, r, 1, label, F_BOLD)
        for j, t in enumerate(turns):
            put(ws, r, 2 + j, sims[key].checks_by[t], fmt=NUM2)
    r += 2
    header(ws, r, 1, ["First Lv6 on the field by end of turn"] + [f"T{t}" for t in turns[:4]])
    for key, label in (("random", "random"), ("first", "first"), ("second", "second")):
        r += 1
        put(ws, r, 1, label, F_BOLD)
        for j, t in enumerate(turns[:4]):
            put(ws, r, 2 + j, sims[key].lv6_by[t], fmt=PCT)
    r += 2
    header(ws, r, 1, ["Bricks and engine usage", "random", "first", "second", "Meaning"])
    metrics = [
        ("No Lv3 digivolve on turn 1", "brick_no_lv3_t1", PCT, "hypergeometric floor is about 10% with the mulligan"),
        ("Mulligan rate", "mulligan_rate", PCT, "policy mulligans any hand without a Lv3"),
        ("No Tamer by end of turn 2", "no_tamer_by_t2", PCT, "engine dead until one shows"),
        ("No Lv4 by end of turn 3", "no_lv4_by_t3", PCT, ""),
        ("Kekkomon attack-time digivolves per game", "tricks", NUM2, ""),
        ("Arts Digivolves per game", "arts", NUM2, ""),
        ("Cougarmon BT25-035 free chains per game", "free_chains", NUM2, ""),
        ("Fuel at the start of the first Lv6 turn", "fuel_at_lv6", NUM1, "the full line needs 4"),
        ("Memory handed over per turn", "memory_given", NUM2, "cap is 3"),
        ("Own security at game end", "security_at_end", NUM1, "how much the policy burned on Barrier-style costs"),
        ("First Lv6 placed on a DUAL Lv5 (no unsuspend)", "lv6_on_dual_base", PCT, "half the Lv5s are DUALs"),
    ]
    for label, attr, fmt, meaning in metrics:
        r += 1
        put(ws, r, 1, label, F_BOLD)
        for j, key in enumerate(("random", "first", "second")):
            put(ws, r, 2 + j, getattr(sims[key], attr), fmt=fmt)
        note(ws, r, 5, meaning)
    r += 2
    header(ws, r, 1, ["Sensitivity (mean kill turn)", "value", "kill by T4", "kill by T5", "note"])
    sens = [
        ("Opponent has a Digimon from its turn 1", sims["opp1"], ""),
        ("Opponent has a Digimon from its turn 2 (default)", sims["edit_base"], "same run as the swap baseline"),
        ("Opponent has a Digimon from its turn 3", sims["opp3"], ""),
        ("Opponent never has a Digimon (T&K memory never fires)", sims["opp99"], ""),
        ("Allow handing over 5 memory instead of 3", sims["give5"], "overspending buys little: the deck is fuel-bound"),
    ]
    for label, s, n in sens:
        r += 1
        put(ws, r, 1, label, F_BOLD)
        put(ws, r, 2, s.kill_mean, fmt=NUM2)
        put(ws, r, 3, s.kill_by[4], fmt=PCT)
        put(ws, r, 4, s.kill_by[5], fmt=PCT)
        note(ws, r, 5, n)
    r += 2
    header(ws, r, 1, ["Card swap: +1 Kyo Sawashiro, -1 e-Pulse", "baseline", "edited", "delta", "noise (1 sd)"])
    a, b = sims["edit_base"], sims["edit_kyo"]
    rows = [
        ("kill by T4", a.kill_by[4], b.kill_by[4], True),
        ("kill by T5", a.kill_by[5], b.kill_by[5], True),
        ("kill turn mean", a.kill_mean, b.kill_mean, False),
        ("Lv6 by T4", a.lv6_by[4], b.lv6_by[4], True),
        ("no Tamer by T2", a.no_tamer_by_t2, b.no_tamer_by_t2, True),
        ("2+ Tamers by end of T3", a.setup[3]["tamers2"], b.setup[3]["tamers2"], True),
        ("tricks per game", a.tricks, b.tricks, False),
    ]
    for label, x, y, is_pct in rows:
        r += 1
        put(ws, r, 1, label, F_BOLD)
        put(ws, r, 2, x, fmt=PCT if is_pct else NUM2)
        put(ws, r, 3, y, fmt=PCT if is_pct else NUM2)
        put(ws, r, 4, f"=C{r}-B{r}", fmt=PCT if is_pct else NUM2)
        if is_pct:
            put(ws, r, 5, f"=SQRT(B{r}*(1-B{r})/{h}+C{r}*(1-C{r})/{h})", fmt=PCT)
        else:
            put(ws, r, 5, None)
    r += 1
    note(ws, r, 1, "A delta smaller than about twice the noise is not evidence of anything.")
    widths(ws, {"A": 52, "B": 13, "C": 13, "D": 13, "E": 13, "F": 11, "G": 11, "H": 15, "I": 15, "J": 11})


def sheet_board(wb, sims):
    ws = wb.create_sheet("Board Setup")
    title(ws, "Board at the end of each turn", "Target going into the Lv6 turn: 2 Tamers and 3+ fuel.")
    rows = [("1+ Tamers", "tamers1", PCT), ("2+ Tamers", "tamers2", PCT), ("fuel (mean)", "fuel_mean", NUM1), ("fuel >= 3", "fuel3", PCT),
            ("Lv4+ out", "top4", PCT), ("Lv5+ out", "top5", PCT), ("Lv6 out", "top6", PCT)]
    r = 4
    for key, label in (("random", "Random first/second"), ("first", "Going first"), ("second", "Going second")):
        s = sims[key]
        turns = sorted(s.setup)
        header(ws, r, 1, [label] + [f"T{t}" for t in turns])
        for lab, k, fmt in rows:
            r += 1
            put(ws, r, 1, lab, F_BOLD)
            for j, t in enumerate(turns):
                put(ws, r, 2 + j, s.setup[t][k], fmt=fmt, fill=FILL_KEY if (lab in ("2+ Tamers", "fuel >= 3") and t == 3) else None)
        r += 2
    note(ws, r, 1, "Reading: fuel lags the curve. You reach Lv6 by paying memory; the 3-4 fuel the multi-attack turn needs arrives a turn later.")
    widths(ws, {"A": 24, "B": 9, "C": 9, "D": 9, "E": 9, "F": 9, "G": 9})


def sheet_setup_math(wb):
    ws = wb.create_sheet("Fuel and Memory")
    title(ws, "Fuel and memory", "What each card gives per turn, and the main combo step by step.")
    header(ws, 4, 1, ["Fuel source", "Fuel per turn per copy", "Condition", "T&K + Kyo", "T&K + Tomoro", "2x T&K + Tomoro", "T&K + Reina (regional)"])
    fuel = [
        ("Tomoro & Kyo (start of main)", 1, "plus 1 on play", 1, 1, 2, 1),
        ("Tomoro Tenma (any Digimon suspends)", 2, "once per cycle; needs an attack", 0, 1, 1, 0),
        ("Kyo Sawashiro (your security removed from)", 1, "Cougarmon ST23-03, Murasamemon BT25-041, their attacks", 1, 0, 0, 0),
        ("Kyo Sawashiro tuck (card from hand)", 1, "costs a card, gives a draw", 1, 0, 0, 0),
        ("Reina Sakuya (any Digimon attacks, theirs too)", 1, "regional list; also grants Blocker", 0, 0, 0, 1),
        ("Reina Sakuya tuck", 1, "", 0, 0, 0, 1),
        ("Gekkomon / Liollmon BT26 moved or played", 1, "one-off per Digimon", 0, 0, 0, 0),
        ("e-Pulse tucks itself", 1, "one-off", 0, 0, 0, 0),
    ]
    r = 5
    first = r
    for row in fuel:
        put(ws, r, 1, row[0], F_BOLD)
        put(ws, r, 2, row[1], F_INPUT)
        note(ws, r, 3, row[2])
        for j in range(4):
            put(ws, r, 4 + j, row[3 + j], F_INPUT)
        r += 1
    last = r - 1
    put(ws, r, 1, "Fuel income per turn", F_BOLD, fill=FILL_KEY)
    for j in range(4):
        col = get_column_letter(4 + j)
        put(ws, r, 4 + j, f"=SUMPRODUCT($B${first}:$B${last},{col}{first}:{col}{last})", F_BOLD, fill=FILL_KEY)
    r += 2
    header(ws, r, 1, ["Memory source", "Memory per turn", "Condition"])
    mem = [("Opponent passes / start of turn", 3, "goldfish assumption"), ("Tomoro & Kyo", 1, "per copy, once they have a Digimon"),
           ("Kyo / Reina / Makoto tuck", 1, "per copy, needs a card in hand"), ("e-Pulse tuck", 1, "the turn after use"),
           ("Glowing Dawn delay", 2, "the turn after use")]
    for lab, v, cnd in mem:
        r += 1
        put(ws, r, 1, lab, F_BOLD)
        put(ws, r, 2, v, F_INPUT)
        note(ws, r, 3, cnd)
    r += 2
    header(ws, r, 1, ["Signature line from a Lv4 (Cougarmon ST23-03 over Liollmon over Kekkomon)", "Memory cost", "Fuel cost", "Memory after", "Fuel after", "Note"])
    r += 1
    put(ws, r, 1, "Start of turn", F_BOLD)
    put(ws, r, 4, 3, F_INPUT)
    put(ws, r, 5, 3, F_INPUT)
    note(ws, r, 6, "blue: edit your starting memory and fuel")
    start = r
    steps = [
        ("Tomoro & Kyo start of main: +1 fuel, +1 memory", -1, -1, ""),
        ("Attack with Cougarmon; Tomoro triggers +2 fuel", 0, -2, "if Tomoro is out"),
        ("Kekkomon trick -2 and Cougarmon's own -2: Murasamemon BT25-041 for 0", 0, 2, "3 - 2 - 2 floors at 0"),
        ("Murasamemon trigger (1 fuel, or your top security): use Atratusmon's Option at 5 - 3", 2, 1, "Arts Digivolve: Murasamemon becomes Atratusmon for free"),
        ("Attack resolves as Atratusmon: 2 checks", 0, 0, "Security Attack +1"),
        ("End of attack: inherited unsuspend", 0, 1, "Murasamemon under the Lv6"),
        ("Second attack: 2 more checks", 0, 0, "4 checks total"),
        ("Close the turn with Kyo (3) or e-Pulse", 3, 0, "end at -1: opponent gets 1"),
    ]
    for lab, m, f, n in steps:
        r += 1
        put(ws, r, 1, lab)
        put(ws, r, 2, m, F_INPUT)
        put(ws, r, 3, f, F_INPUT)
        put(ws, r, 4, f"=D{r - 1}-B{r}")
        put(ws, r, 5, f"=E{r - 1}-C{r}")
        note(ws, r, 6, n)
    r += 1
    note(ws, r, 1, f"Negative memory after a step means the turn ends there (opponent's side). Fuel after must stay >= 0 for the line to be legal; edit D{start}:E{start} to test a start.")
    widths(ws, {"A": 72, "B": 16, "C": 46, "D": 14, "E": 14, "F": 20, "G": 20})


# Their 50 cards as security, counted from the regional lists (card DP from data/cardtext.json
# and goldfish.cards). Lives on the Matchups sheet; the Calculator reads the same cells.
MATCHUPS = ["Jupitermon", "TS Mervamon (Homeros control)", "Toho Braves (TB)", "DATA SQUAD Ravemon (Rose Rave)"]
CRIMSON_BLAZE = "BT8-097"
CARDTEXT = json.loads((ROOT / "data" / "cardtext.json").read_text())


def security_composition(cards):
    """(Counter of Digimon DP, tamers, crimson, other options) for one 50-card list."""
    dps, tamers, crimson, other = Counter(), 0, 0, 0
    for n, _, cid, lvl in cards:
        if lvl == "Egg":
            continue
        if cid in CARDTEXT:
            kind, dp = CARDTEXT[cid].get("cardtype"), CARDTEXT[cid].get("dp")
        elif cid in CARDS:
            kind, dp = CARDS[cid].kind.value.capitalize(), CARDS[cid].dp
        else:
            raise KeyError(f"no card data for {cid}; run tools/fetch_card_text.py {cid} --json data/cardtext.json")
        if dp:
            dps[int(dp)] += n
        elif kind == "Tamer":
            tamers += n
        elif cid == CRIMSON_BLAZE:
            crimson += n
        else:
            other += n
    return dps, tamers, crimson, other


SEC_DETAIL = [
    "12000+: Jupitermon x4, Wide Plasment x1, Wrath Mode x3, Chronomon x1, Dianamon x1. 8000: Aegiochusmon x8. 5000: Aegiomon x8. Blinding Ray has no [Security] effect.",
    "12000+: Mervamon x4, Minervamon x3, Bacchusmon x3, Ceresmon x2, Dianamon x1, Wide Plasment x3, Chaosmon x1 (17!). 8000: Aegiochusmon Holy x2. 6000: Sirenmon x4; 5000: Aegiomon x3. Homeros x4 among the Tamers. Central Town x2 plays a Lv4 free.",
    "12000+: Amaterasumon x3, Kaguyamon x3, Ryugumon x1, Susanoomon x3. 7000: Shishimamon x4, Karakurumon x2, MarineBullmon x1. 5000-6000: Kokeshimon x4, Manekimon x4, Musyamon, Shellmon. Genshi x3 plays a <=5-cost card free; Sanmyojin x4 draws 2; Decree x1.",
    "12000+: Ravemon x4, Rosemon x2, Rosemon BM x3, ShineGreymon BM x1. 7000: Crowmon x4, Lilamon x3. 5000-6000: GeoGreymon x4, Peckmon x4. DNA Charge x4 plays a <=4-cost card free.",
]
MATCHUP_TEXT = {
    "Removal on their turn": [
        "Jupitermon -13000 DP on play / digivolve (kills anything we own); Wrath Mode -15000 whenever security is removed; Dianamon deletes an unsuspended Digimon; Chronomon deletes <=12000; Crimson Blaze wipes <=6000.",
        "Minervamon De-Digivolve 1 per Digimon they have; Bacchusmon deletes our lowest DP; Jupitermon Option deletes all our lowest DP; Chaosmon lock (no unsuspend, no On Play); Mervamon -4000 per Iliad card.",
        "Execute: end-of-turn attacks into our UNSUSPENDED Digimon (Susanoomon 16000 Rush); Kokeshimon deletes <=Lv4; Musyamon <=6000; Amaterasumon deletes lowest DP; Susanoomon places our Digimon into THEIR security.",
        "Peckmon x4 and Crowmon x4 delete <=Lv4 (our Kekkomon base dies on sight); Ravemon deletes our HIGHEST DP, repeatable at end of attack for 2 fuel; Lilamon and Rosemon BM suspend-lock (no unsuspend, no digivolve).",
    ],
    "Our answer": [
        "Atratusmon's immunity (their Digimon effects, lasts through their turn), made BEFORE the first check of our turn. Habakirimon's Option side (-8000, -5000 to all) is the board wipe. Neither stops Options (Wide Plasment, Crimson Blaze).",
        "Pass exactly 4 so every Homeros suspends itself. Suspend Minervamon instead of deleting it; remove Sirenmon first; Armalizamon suspends a blocker, Eclipse Impact bottom-decks the biggest. Mervamon's -DP and Wide Plasment get around immunity.",
        "Atratusmon's immunity blanks Amaterasumon, Shishimamon and Susanoomon (-DP and security placement). Barrier answers Execute. Eclipse Impact bottom-decks Kaguyamon / Ryugumon (no Retaliation, no On Deletion). Keep the Lv4 in breeding.",
        "Immunity answers Ravemon, Crowmon and Rosemon once Atratusmon has digivolved or attacked. The Rosemon Tamer lock is the real threat, not hand rip: 2 Tamers early and a second attacker. Raise to Lv5 in breeding.",
    ],
    "Blockers they present (after their T2 / T3)": ["1 / 2-3 (Wrath Mode; Aegiochusmon Blue)", "2 / 4-5 (every Iliad Digimon under Minervamon or Mervamon, with Reboot)", "1, or 4 with Kaguyamon / 3-4 (Puppets get Retaliation)", "1 / 1 (Peckmon; no real blocker plan)"],
    "Their clock (research; their own turns)": ["Burst T2 (Lv5-Lv7, 2 removals, 4 checks); kill T3 (n=1)", "Minervamon T2, Mervamon T3; attacks only when safe; T5+ or a time draw", "Burst T2 (Shishimamon, Execute, Lv6); Susanoomon T3; fastest kill T3", "Combo T2 (2-4 removals, hand strip), T3 'more likely'; Rosemon lock T2-T3; kill T4-T5 (very low confidence)"],
    "Lv6 to aim for": ["Atratusmon; Habakirimon for its Option side", "Atratusmon", "Atratusmon; Habakirimon second", "Atratusmon"],
    "Breeding plan": ["Evolve in the back; nothing we leave out survives their T2", "Move Lv3 on T2 is fine", "Raise to Lv5 in breeding, move T3", "Raise to Lv5 in breeding, move T3"],
    "Attack pacing": ["Don't chip, and end the turn with them at 0 security: each check trashes ours and fires Holy / Wrath", "Build freely (their T2 had no checks into us), but expect Mervamon on their T3. Don't poke (20% free Tamer, 8% free Homeros)", "Take all 5 in one turn or steal the turn: chip feeds Susanoomon", "Batch attacks; a 1-check poke feeds a free Tamer 20% of the time"],
}


def sheet_matchups(wb):
    """Returns the row numbers other sheets link to: win rates and per-check probabilities."""
    ws = wb.create_sheet("Matchups")
    title(ws, "Matchups", "Real win rates, and what their security does to your attackers.")
    header(ws, 4, 1, ["Real results, Glowing Dawn's side (DigiLab)", "GD wins", "GD losses", "Ties", "Win rate", "Likely range: low", "Likely range: high", "Their side", "Note"])
    z = "1.96"  # 95% Wilson score interval
    wr_rows = {}
    r = 4
    for name, w, lo, t, theirs, why in WINRATES:
        r += 1
        put(ws, r, 1, name, F_BOLD)
        put(ws, r, 2, w, F_INPUT)
        put(ws, r, 3, lo, F_INPUT)
        put(ws, r, 4, t, F_INPUT)
        put(ws, r, 5, f"=B{r}/(B{r}+C{r})", fmt=PCT, fill=FILL_KEY)
        n, p = f"(B{r}+C{r})", f"E{r}"
        centre = f"({p}+{z}^2/(2*{n}))/(1+{z}^2/{n})"
        half = f"{z}*SQRT({p}*(1-{p})/{n}+{z}^2/(4*{n}^2))/(1+{z}^2/{n})"
        put(ws, r, 6, f"={centre}-{half}", fmt=PCT)
        put(ws, r, 7, f"={centre}+{half}", fmt=PCT)
        put(ws, r, 8, theirs, wrap=True)
        put(ws, r, 9, why, wrap=True)
        wr_rows[name] = r
    r += 1
    note(ws, r, 1, "Likely range = 95% confidence (Wilson). Only 25 games vs TS Mervamon, so its true rate could be anywhere from about 6% to 35%.")
    for event, text, tag in EVENT_RESULTS:
        r += 1
        put(ws, r, 1, event, F_BOLD)
        put(ws, r, 2, text, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        put_tag(ws, r, 9, tag)
    r += 2
    header(ws, r, 1, ["Their 50 cards as security: Digimon by DP, then other cards (counted from the lists)"] + MATCHUPS)
    hdr_row = r
    r += 1
    comps = [security_composition(cards) for head, cards in META_LISTS.items() if not head.startswith("Glowing Dawn")]
    assert len(comps) == len(MATCHUPS)
    first = r
    for dp in sorted({dp for c in comps for dp in c[0]}, reverse=True):
        put(ws, r, 1, dp, F_BOLD, fmt='#,##0" DP"', align="left")
        for j, comp in enumerate(comps):
            put(ws, r, 2 + j, comp[0].get(dp, 0), F_INPUT)
        r += 1
    special = {}
    for label, idx in (("Tamers ([Security] play free)", 1), ("Crimson Blaze ([Security] deletes all our <= 6000 DP)", 2),
                       ("Other Options (free plays, draws, placements)", 3)):
        put(ws, r, 1, label, F_BOLD, wrap=True)
        for j, comp in enumerate(comps):
            put(ws, r, 2 + j, comp[idx], F_INPUT)
        special[idx] = r
        r += 1
    last = r - 1
    put(ws, r, 1, "Total (must be 50)", F_BOLD)
    for j in range(4):
        col = get_column_letter(2 + j)
        put(ws, r, 2 + j, f"=SUM({col}{first}:{col}{last})", F_BOLD)
    total = r
    r += 1
    for j, d in enumerate(SEC_DETAIL):
        put(ws, r + j, 1, f"Detail: {MATCHUPS[j]}", F_NOTE)
        put(ws, r + j, 2, d, F_NOTE, wrap=True)
        ws.merge_cells(start_row=r + j, start_column=2, end_row=r + j, end_column=5)
    r += 5
    header(ws, r, 1, ["Per security check we make"] + MATCHUPS)
    dp_range = f"$A${first}:$A${last}"
    kills = lambda dp: f'=SUMIF({dp_range},">={dp}",{{c}}${first}:{{c}}${last})/{{c}}{total}'  # noqa: E731
    probs = [
        ("Kills a 12000 attacker (Atratusmon, Habakirimon)", kills(12000)),
        ("Kills a 7000 Lv5 attacker", kills(7000)),
        ("Kills a 4000 Lv4 attacker", kills(4000)),
        ("Gives them a free Tamer", f"={{c}}{special[1]}/{{c}}{total}"),
        ("Wipes our <= 6000 DP Digimon", f"={{c}}{special[2]}/{{c}}{total}"),
        ("Free play / resource for them", f"={{c}}{special[3]}/{{c}}{total}"),
    ]
    prob_rows = {}
    for label, tmpl in probs:
        r += 1
        put(ws, r, 1, label, F_BOLD)
        for j in range(4):
            col = get_column_letter(2 + j)
            put(ws, r, 2 + j, tmpl.format(c=col), fmt=PCT)
        prob_rows[label] = r
    r += 2
    put(ws, r, 1, "Checks in a multi-attack turn", F_BOLD)
    put(ws, r, 2, 4, F_INPUT)
    note(ws, r, 3, "Atratusmon twice = 4; Habakirimon three times = 3")
    nrow = r
    r += 1
    header(ws, r, 1, ["P(at least one attacker-killer flips in that turn)"] + MATCHUPS)
    for label, pr in (("12000 attacker", prob_rows["Kills a 12000 attacker (Atratusmon, Habakirimon)"]), ("7000 Lv5 attacker", prob_rows["Kills a 7000 Lv5 attacker"])):
        r += 1
        put(ws, r, 1, label, F_BOLD)
        for j in range(4):
            col = get_column_letter(2 + j)
            put(ws, r, 2 + j, f"=1-(1-{col}{pr})^$B${nrow}", fmt=PCT, fill=FILL_KEY if label.startswith("12000") else None)
    r += 1
    note(ws, r, 1, "Without Barrier under the Lv6 the multi-attack turn ends in a dead Lv6 more often than not. With Barrier each hit costs a security card instead.")
    r += 2
    header(ws, r, 1, ["Their turn, and ours"] + MATCHUPS)
    for label, cells in MATCHUP_TEXT.items():
        r += 1
        put(ws, r, 1, label, F_BOLD)
        for j, txt in enumerate(cells):
            put(ws, r, 2 + j, txt, wrap=True)
    widths(ws, {"A": 58, "B": 42, "C": 42, "D": 42, "E": 42, "F": 12, "G": 12, "H": 34, "I": 50})
    return {"winrate": wr_rows, "prob": prob_rows, "sec_first": first, "sec_last": last, "sec_total": total,
            "sec_special": special, "header": hdr_row}


META_LISTS = {
    "Jupitermon - 1st of 181, Fish, Regionals @ NoHeroes Dusseldorf 2026-10-04 (7-1)": [
        (4, "Tsunomon", "BT24-003", "Egg"), (4, "Elecmon", "BT25-030", "Lv3"), (4, "Elecmon", "BT24-031", "Lv3"), (2, "Lunamon", "BT25-022", "Lv3"),
        (1, "Liollmon", "BT26-025", "Lv3"), (4, "Aegiomon", "BT25-033", "Lv4"), (4, "Aegiomon", "BT24-034", "Lv4"), (2, "Aegiochusmon", "P-213", "Lv5"),
        (2, "Aegiochusmon: Blue", "BT25-025", "Lv5"), (4, "Aegiochusmon: Holy", "BT26-029", "Lv5"), (1, "Chronomon: Holy Mode", "BT26-016", "Lv6"),
        (1, "Dianamon", "BT25-028", "Lv6"), (4, "Jupitermon", "BT24-101", "Lv6"), (1, "Jupitermon // Wide Plasment", "BT26-033", "Lv6 DUAL"),
        (3, "Jupitermon: Wrath Mode", "BT26-103", "Lv7"), (2, "Hiroko Sagisaka", "BT24-083", "Tamer"), (1, "Hiroko Sagisaka", "P-210", "Tamer"),
        (4, "Inori Misono", "BT24-084", "Tamer"), (1, "Shota Kuroi", "BT26-092", "Tamer"), (1, "Toya Kuga", "BT26-087", "Tamer"),
        (1, "Kosuke Misono", "BT26-096", "Tamer"), (1, "Blinding Ray", "BT4-104", "Option"), (2, "Crimson Blaze", "BT8-097", "Option"),
    ],
    "TS Mervamon (Homeros control) - 5th of 181, UnkindledOne, Dusseldorf (6-1-1)": [
        (4, "Wanyamon", "BT24-004", "Egg"), (3, "Coronamon", "BT25-008", "Lv3"), (4, "Elecmon", "BT24-031", "Lv3"), (1, "Floramon", "BT25-047", "Lv3"),
        (4, "Tapirmon", "BT24-043", "Lv3"), (3, "Aegiomon", "BT24-034", "Lv4"), (2, "Aegiochusmon: Holy", "BT26-029", "Lv5"), (4, "Sirenmon", "BT25-039", "Lv5"),
        (3, "Bacchusmon", "BT25-077", "Lv6"), (1, "Ceresmon", "BT25-059", "Lv6"), (1, "Ceresmon // Famis", "BT26-032", "Lv6 DUAL"), (1, "Dianamon", "BT25-028", "Lv6"),
        (3, "Jupitermon // Wide Plasment", "BT26-033", "Lv6 DUAL"), (4, "Mervamon", "BT26-081", "Lv6"), (3, "Minervamon", "BT24-041", "Lv6"),
        (1, "Chaosmon: Valdur Arm", "BT20-037", "Lv7"), (3, "Kanan Yuki", "BT26-090", "Tamer"), (1, "Shota Kuroi", "BT26-092", "Tamer"),
        (1, "Toya Kuga", "BT26-087", "Tamer"), (1, "Dan Yuki & Kanan Yuki", "BT24-085", "Tamer"), (4, "Homeros", "BT24-102", "Tamer"),
        (2, "Central Town: Throne Room", "BT24-094", "Option"),
    ],
    "Toho Braves (TB) - 10th of 181, Vaelthas, Dusseldorf (6-2); archetype also won Sao Paulo (118 players, 2026-09-19)": [
        (4, "Onibimon", "EX12-004", "Egg"), (2, "Gasamon", "EX12-020", "Lv3"), (4, "Hanimon", "EX12-061", "Lv3"), (1, "Kotemon", "BT26-008", "Lv3"),
        (4, "Wankomon", "EX12-009", "Lv3"), (4, "Kokeshimon", "EX12-062", "Lv4"), (4, "Manekimon", "BT26-012", "Lv4"), (1, "Musyamon", "BT26-013", "Lv4"),
        (1, "Shellmon", "EX12-026", "Lv4"), (2, "Karakurumon", "EX12-063", "Lv5"), (1, "MarineBullmon", "EX12-031", "Lv5"), (4, "Shishimamon", "EX12-046", "Lv5"),
        (3, "Amaterasumon", "EX12-047", "Lv6"), (3, "Kaguyamon", "EX12-065", "Lv6"), (1, "Ryugumon", "EX12-036", "Lv6"), (3, "Susanoomon", "EX12-076", "Lv7"),
        (2, "Kunlun", "BT26-104", "Tamer"), (3, "Genshi Continent & Ashino Island", "EX12-074", "Option"), (1, "Kunlun's Imperial Decree", "EX12-075", "Option"),
        (4, "Sanmyojin Arrival", "EX12-070", "Option"), (2, "Crimson Blaze", "BT8-097", "Option"),
    ],
    "DATA SQUAD Ravemon (Rose Rave) - 16th of 181, Rayquon, Dusseldorf (5-1-2)": [
        (4, "Pinamon", "BT26-005", "Egg"), (4, "Agumon", "ST24-04", "Lv3"), (1, "Falcomon", "BT26-065", "Lv3"), (2, "Falcomon", "ST24-12", "Lv3"),
        (4, "Lalamon", "BT26-036", "Lv3"), (4, "GeoGreymon", "ST24-05", "Lv4"), (4, "Peckmon", "BT26-072", "Lv4"), (4, "Crowmon", "BT26-076", "Lv5"),
        (3, "Lilamon", "ST24-10", "Lv5"), (4, "Ravemon", "BT26-082", "Lv6"), (2, "Rosemon", "BT26-049", "Lv6"), (3, "Rosemon: Burst Mode", "BT26-050", "Lv7 DUAL"),
        (1, "ShineGreymon: Burst Mode", "BT25-104", "Lv7 DUAL"), (4, "Keenan Crier", "BT26-094", "Tamer"), (4, "Yoshino Fujieda", "BT26-091", "Tamer"),
        (2, "Yoshino Fujieda & Keenan Crier", "ST24-14", "Tamer"), (4, "DNA Charge", "ST24-15", "Option"),
    ],
    "Glowing Dawn - 4th of 181, Quang-Minh [Gere Gaming], Dusseldorf (6-0-2)": [
        (4, "Kekkomon", "ST23-01", "Egg"), (3, "Chiropmon", "ST23-12", "Lv3"), (4, "Gekkomon", "ST23-06", "Lv3"), (2, "Liollmon", "BT25-032", "Lv3"),
        (4, "Ukkomon", "BT16-082", "Lv3"), (3, "Armalizamon", "BT25-049", "Lv4"), (4, "Cougarmon", "ST23-03", "Lv4"), (1, "Cougarmon", "BT25-035", "Lv4"),
        (3, "Monarchlizamon", "BT25-057", "Lv5 DUAL"), (2, "Murasamemon", "BT25-041", "Lv5"), (4, "Murasamemon", "ST23-04", "Lv5"),
        (4, "Atratusmon", "ST23-09", "Lv6 DUAL"), (2, "Habakirimon", "BT25-043", "Lv6 DUAL"), (1, "Makoto Kuonji", "BT26-095", "Tamer"),
        (3, "Reina Sakuya", "BT26-093", "Tamer"), (2, "Reina Sakuya & Makoto Kuonji", "ST23-14", "Tamer"),
        (4, "Tomoro Tenma & Kyo Sawashiro", "ST23-13", "Tamer"), (4, "e-Pulse", "ST23-15", "Option"),
    ],
}
META_LINKS = {
    "Jupitermon": "https://digimon.digilab.cards/decklist/116163",
    "TS Mervamon": "https://digimon.digilab.cards/decklist/116167",
    "Toho Braves": "https://digimon.digilab.cards/decklist/116172",
    "DATA SQUAD": "https://digimon.digilab.cards/decklist/116178",
    "Glowing Dawn": "https://digimon.digilab.cards/decklist/116166",
}


def sheet_meta_lists(wb):
    ws = wb.create_sheet("Meta Decklists")
    title(ws, "Meta decklists", "Regional top-cut lists from DigiLab (links on the Sources sheet).")
    r = 4
    for head, cards in META_LISTS.items():
        header(ws, r, 1, [head, "", "", ""])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
        header(ws, r, 1, ["Count", "Card", "ID", "Level / type"], fill=FILL_SUB, font=F_BOLD)
        r += 1
        first = r
        for n, name, cid, lvl in cards:
            put(ws, r, 1, n)
            put(ws, r, 2, name)
            put(ws, r, 3, cid)
            put(ws, r, 4, lvl)
            r += 1
        last = r - 1
        put(ws, r, 1, f'=SUMIF($D${first}:$D${last},"<>Egg",$A${first}:$A${last})', F_BOLD)
        put(ws, r, 2, "main deck total (50)", F_BOLD)
        put(ws, r, 3, f'=SUMIF($D${first}:$D${last},"Egg",$A${first}:$A${last})', F_BOLD)
        put(ws, r, 4, "eggs", F_BOLD)
        r += 3
    widths(ws, {"A": 8, "B": 34, "C": 11, "D": 12})


EXTRA_CARDS = {"BT16-082": ("Ukkomon", "Lv3"), "BT26-093": ("Reina Sakuya", "Tamer"), "ST23-14": ("Reina Sakuya & Makoto Kuonji", "Tamer"), "BT26-095": ("Makoto Kuonji", "Tamer")}
EXTRA_NOTES = {
    "BT16-082": "On move from breeding: reveal 3, add a Digimon or Tamer, then re-hatch. Not Glowing Dawn.",
    "BT26-093": "Tuck -> draw 1, +1 memory. When ANY Digimon attacks: +1 fuel and a BEATBREAK Digimon gains Blocker + Collision.",
    "ST23-14": "T&K clone (fuel + memory); when its fuel is trashed, a Glowing Dawn Digimon gains Jamming.",
    "BT26-095": "Tuck -> draw 1, +1 memory. When any Digimon is deleted: draw 1 trash 1, recycle a card from trash as fuel.",
}


def sheet_comparison(wb, deck):
    ws = wb.create_sheet("Deck Comparison")
    title(ws, "Your list vs the 4th-place Dusseldorf list", "Positive difference = they run more copies.")
    mine = deck.counts()
    theirs = {cid: n for n, _, cid, _ in META_LISTS[next(k for k in META_LISTS if k.startswith("Glowing Dawn"))]}
    ids = sorted(set(mine) | set(theirs), key=lambda cid: ((CARDS[cid].level or (9 if CARDS[cid].kind.value != "tamer" else 8)) if cid in CARDS else {"Lv3": 3, "Tamer": 8}[EXTRA_CARDS[cid][1]], cid))
    header(ws, 4, 1, ["Card", "ID", "Level / type", "Yours", "Regional 4th", "Delta", "Why it matters"])
    r = 5
    first = r
    for cid in ids:
        if cid in CARDS:
            c = CARDS[cid]
            if c.kind.value == "egg":
                lvl = "Egg"
            elif c.level:
                lvl = f"Lv{c.level}" + (" DUAL" if c.kind.value == "dual" else "")
            else:
                lvl = c.kind.value.capitalize()
            name = c.name
            why = ROLE_NOTES.get(cid, "")
        else:
            name, lvl = EXTRA_CARDS[cid]
            why = EXTRA_NOTES[cid]
        put(ws, r, 1, name)
        put(ws, r, 2, cid)
        put(ws, r, 3, lvl)
        put(ws, r, 4, mine.get(cid, 0), F_INPUT)
        put(ws, r, 5, theirs.get(cid, 0), F_INPUT)
        put(ws, r, 6, f"=E{r}-D{r}")
        put(ws, r, 7, why, wrap=True)
        r += 1
    last = r - 1
    r += 1
    header(ws, r, 1, ["Totals", "", "", "Yours", "Regional 4th", "Delta"], fill=FILL_SUB, font=F_BOLD)
    groups = [("Lv3", "Lv3"), ("Lv4", "Lv4"), ("Lv5 (all)", "Lv5*"), ("Lv5 non-DUAL", "Lv5"), ("Lv6 (all)", "Lv6*"), ("Tamers", "Tamer"), ("Options", "Option"), ("Eggs", "Egg")]
    for label, crit in groups:
        r += 1
        put(ws, r, 1, label, F_BOLD)
        for col in ("D", "E"):
            put(ws, r, 4 if col == "D" else 5, f'=SUMIF($C${first}:$C${last},"{crit}",{col}${first}:{col}${last})')
        put(ws, r, 6, f"=E{r}-D{r}")
    r += 2
    put(ws, r, 1, "Reading", F_BOLD)
    put(ws, r, 2, "Every difference trades speed for survivability: Barrier bases (4 Cougarmon ST23-03), a defensive fuel engine (Reina), the -5000 DP Lv5 (ST23-04), 4 Atratusmon blockers, and recursion (3 Chiropmon). Consistent with a field where the Lv4 dies on sight and the Lv6 eats -13000.", wrap=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 45
    widths(ws, {"A": 32, "B": 11, "C": 12, "D": 8, "E": 13, "F": 8, "G": 90})



# ---------------------------------------------------------------------- calculator
# Input cells on the Calculator sheet (tools/check_workbook.py reads these too).
CALC_IN = {"matchup": "C5", "their_security": "C6", "stoppers": "C7", "stopper_dp": "C8", "minus": "C9",
           "your_security": "C10", "first_check_trash": "C11"}
CALC_ATTACKER_ROW = 15  # attackers on rows 15-17: B name, C DP, D checks, E attacks, F Barrier, G immune, H Piercing
CALC_ATTACKER_COLS = {"name": "B", "dp": "C", "checks": "D", "attacks": "E", "barrier": "F", "immune": "G", "piercing": "H"}
CALC_DEFAULTS = {"matchup": MATCHUPS[0], "their_security": 2, "stoppers": 1, "stopper_dp": 16000, "minus": 5000,
                 "your_security": 5, "first_check_trash": 1}
CALC_DEFAULT_ATTACKERS = [
    ("Atratusmon (attacks first)", 12000, 2, 2, "N", "Y", "N"),
    ("Murasamemon ST23-04", 7000, 1, 1, "N", "N", "N"),
    ("(empty: set Attacks to 0)", 0, 1, 0, "N", "N", "N"),
]
CALC_RESULTS = [  # (label, Result field, format)
    ("P(you win this turn)", "p_win", PCT),
    ("P(they end at 0 security, no win yet)", "p_zero", PCT),
    ("Their security left (average)", "security_left", NUM2),
    ("Checks you make (average)", "checks", NUM2),
    ("Free Tamers they flip (average)", "free_tamers", NUM2),
    ("Free plays / resources they flip (average)", "free_plays", NUM2),
    ("Your attackers lost (average)", "attackers_lost", NUM2),
    ("Your security paid to Barrier (average)", "barrier_paid", NUM2),
    ("Your security after the turn (average)", "your_security_after", NUM2),
]
CALC_RESULT_ROW = 21  # first result row; columns C (early), D (lethal), E (plan for)
ENGINE = "Calc Engine"
ENG = f"'{ENGINE}'"


def sheet_calculator(wb, links):
    ws = wb.create_sheet("Calculator")
    title(ws, "Turn calculator",
          "Fill in the blue cells. Odds come from their real decklist and the official rules.")
    header(ws, 4, 2, ["Situation", "Value", "How to read it", "", "", "", ""])
    ws.merge_cells("D4:H4")
    situation = [
        ("matchup", "Matchup", "picks their 50-card list (Matchups sheet)"),
        ("their_security", "Their security now", "0-7"),
        ("stoppers", "Attacks they can stop this turn", "unsuspended Blockers + Sirenmon's inherited redirect + Aegiochusmon: Blue's inherited unsuspend (re-readies a blocker); 0-3"),
        ("stopper_dp", "DP of what stops you", "the blocker you would battle (Wrath Mode 16000, Aegiochusmon Blue 8000, Mervamon 15000, Kaguyamon Puppets, Peckmon 5000)"),
        ("minus", "-DP their check triggers give your non-immune attackers", "Aegiochusmon: Holy on board = 5000 (fires on your first check); add 15000 for Wrath Mode only if it targets that attacker"),
        ("your_security", "Your security now", "Barrier and Habakirimon's protection spend it"),
        ("first_check_trash", "Your security they trash on your first check", "1 if Jupitermon is on their board, else 0"),
    ]
    for i, (key, label, how) in enumerate(situation):
        r = 5 + i
        put(ws, r, 2, label, F_BOLD, wrap=True)
        put(ws, r, 3, CALC_DEFAULTS[key], F_INPUT, fill=FILL_KEY if key == "their_security" else None, align="left")
        put(ws, r, 4, how, F_NOTE, wrap=True)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    header(ws, 14, 2, ["Your attackers, in attack order", "DP now", "Checks per attack (Security A.)", "Attacks this turn (2 = it unsuspends)",
                       "Barrier under it (Y/N)", "Immune to their Digimon effects (Y/N)", "Piercing (Y/N)",
                       "DP in security battles", "Dies per check", "Barrier pays per check"])
    for i, row in enumerate(CALC_DEFAULT_ATTACKERS):
        r = CALC_ATTACKER_ROW + i
        for j, v in enumerate(row):
            put(ws, r, 2 + j, v, F_INPUT, align="left")
        for j, field in enumerate(("checkdp", "die", "pay")):
            put(ws, r, 9 + j, f'=IF(E{r}=0,"",{ENG}!{get_column_letter(3 + i)}{ENGINE_ATT_ROWS[field]})', F_LINK,
                fmt=None if field == "checkdp" else PCT, align="right")
    note(ws, 18, 2, "Inherited effects: Barrier comes from Liollmon BT25 / Cougarmon under the stack, Piercing from Gekkomon / Armalizamon, "
                    "the second attack from a non-DUAL Lv5's end-of-attack unsuspend (DUAL Lv5s give none).")
    note(ws, 19, 2, "Immune = Atratusmon once it has digivolved or attacked this turn: their -DP triggers can't touch it, so let it attack first.")
    header(ws, CALC_RESULT_ROW - 1, 2, ["Result", "They block early", "They save blocks for lethal", "Plan for (lower win chance)"])
    for i, (label, field, fmt) in enumerate(CALC_RESULTS):
        r = CALC_RESULT_ROW + i
        put(ws, r, 2, label, F_BOLD)
        for j in range(len(gc.POLICIES)):
            put(ws, r, 3 + j, f"={ENG}!{ENGINE_OUT[(gc.POLICIES[j], field)]}", F_LINK, fmt=fmt)
        put(ws, r, 5, f"=IF($C${CALC_RESULT_ROW}<=$D${CALC_RESULT_ROW},C{r},D{r})", F_BOLD, fmt=fmt,
            fill=FILL_KEY if field == "p_win" else None)
    r = CALC_RESULT_ROW + len(CALC_RESULTS)
    put(ws, r, 2, "Barrier check", F_BOLD)
    put(ws, r, 3, f'=IF(E{r - 2}>C10,"You may run out of security for Barrier: the odds above are optimistic","OK")', wrap=True)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    r += 2
    notes = [
        "Win = an attack connects while they have 0 security. A check that removes their last card is not a win; the next unblocked attack is.",
        "A revealed security Digimon with DP >= your attacker deletes it (ties delete both); Barrier trashes your top security instead. Crimson Blaze deletes a <= 6000 DP attacker, Barrier or not.",
        "Each check is drawn independently from their 50-card list. Real security is 5 specific cards, so treat results as the average over all the ways their security could be stacked.",
        "'Block early' = they stop your first attacks; 'save for lethal' = they stop only the attack that would win. Real opponents sit in between, so plan for the lower one.",
        "Not modeled: Wrath Mode's -15000 picking a target (enter it in the -DP cell or lower that attacker's DP), security effects that play blockers (Toho's Island), Habakirimon's security trash.",
    ]
    for t in notes:
        put(ws, r, 2, t, F_NOTE, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
        r += 1

    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"' + ",".join(MATCHUPS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(CALC_IN["matchup"])
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    ws.add_data_validation(yn)
    yn.add(f"F{CALC_ATTACKER_ROW}:H{CALC_ATTACKER_ROW + 2}")
    for cell, lo, hi in ((CALC_IN["their_security"], 0, gc.S_MAX), (CALC_IN["stoppers"], 0, gc.B_MAX),
                         (CALC_IN["first_check_trash"], 0, 1), (f"D{CALC_ATTACKER_ROW}:D{CALC_ATTACKER_ROW + 2}", 1, gc.MAX_CHECKS),
                         (f"E{CALC_ATTACKER_ROW}:E{CALC_ATTACKER_ROW + 2}", 0, gc.MAX_ATTACKS)):
        v = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi))
        ws.add_data_validation(v)
        v.add(cell)
    widths(ws, {"A": 2, "B": 34, "C": 16, "D": 18, "E": 18, "F": 13, "G": 15, "H": 11, "I": 12, "J": 10, "K": 11})


# Engine layout, filled in by sheet_calc_engine (Calculator links to these addresses).
ENGINE_ATT_ROWS = {}
ENGINE_OUT = {}


def sheet_calc_engine(wb, links):
    """Live formulas for goldfish.calc: one column per step of the turn, one row per state."""
    ws = wb.create_sheet(ENGINE)
    title(ws, "Calculator engine (generated; edit goldfish/calc.py, not these cells)",
          "Rows = states (their security s, your current attacker dead / idle / checking, stoppers left b). Columns = steps of the turn. Each cell = probability of that state after that step.")
    calc = "Calculator!"
    m = links
    # globals
    g = {}
    rows = [
        ("idx", f"=MATCH({calc}${CALC_IN['matchup'][0]}${CALC_IN['matchup'][1:]},Matchups!$B${m['header']}:$E${m['header']},0)", "matchup column"),
        ("s0", f"=MIN(MAX({calc}{CALC_IN['their_security']},0),{gc.S_MAX})", "their security, capped"),
        ("b0", f"=MIN(MAX({calc}{CALC_IN['stoppers']},0),{gc.B_MAX})", "stoppers, capped"),
        ("y", f"={calc}{CALC_IN['stopper_dp']}", "stopper DP"),
        ("minus", f"={calc}{CALC_IN['minus']}", "-DP from their check triggers"),
        ("total", None, "cards in their list"),
        ("pt", None, "P(Tamer) per check"),
        ("po", None, "P(other free play) per check"),
        ("pc", None, "P(Crimson Blaze) per check"),
    ]
    r = 4
    header(ws, r, 1, ["Global", "Value", "Meaning"], fill=FILL_SUB, font=F_BOLD)
    for key, f, meaning in rows:
        r += 1
        g[key] = f"$B${r}"
        put(ws, r, 1, key, F_BOLD)
        put(ws, r, 2, f)
        note(ws, r, 3, meaning)
    # selected matchup's security composition
    r += 2
    header(ws, r, 1, ["Their DP / card", "Count (selected list)"], fill=FILL_SUB, font=F_BOLD)
    sec_first = r + 1
    for mr in range(m["sec_first"], m["sec_last"] + 1):
        r += 1
        put(ws, r, 1, f"=IF(ISNUMBER(Matchups!$A{mr}),Matchups!$A{mr},\"\")")
        put(ws, r, 2, f"=INDEX(Matchups!$B{mr}:$E{mr},{g['idx']})")
    sec_last = r
    dp_rng, n_rng = f"$A${sec_first}:$A${sec_last}", f"$B${sec_first}:$B${sec_last}"
    sel = lambda mr: f"INDEX(Matchups!$B{mr}:$E{mr},{g['idx']})"  # noqa: E731
    ws[g["total"].replace("$", "")] = f"=SUM({n_rng})"
    ws[g["pt"].replace("$", "")] = f"={sel(m['sec_special'][1])}/{g['total']}"
    ws[g["po"].replace("$", "")] = f"={sel(m['sec_special'][3])}/{g['total']}"
    ws[g["pc"].replace("$", "")] = f"={sel(m['sec_special'][2])}/{g['total']}"
    # per-attacker helpers (columns C-E)
    r += 2
    header(ws, r, 1, ["Attacker helper", ""] + [f"Attacker {i + 1}" for i in range(gc.MAX_ATTACKERS)], fill=FILL_SUB, font=F_BOLD)
    ar = CALC_ATTACKER_ROW
    yes = lambda col: f'IF(UPPER({calc}{col}{{row}})="Y",1,0)'  # noqa: E731
    helpers = [
        ("dp", f"=N({calc}C{{row}})"),
        ("checks", f"=N({calc}D{{row}})"),
        ("attacks", f"=N({calc}E{{row}})"),
        ("barrier", "=" + yes("F")),
        ("immune", "=" + yes("G")),
        ("piercing", "=" + yes("H")),
        ("checkdp", "={c}{dp}-IF({c}{immune}=1,0,{minus})"),
        ("battleloss", f'=IF({{c}}{{checkdp}}<=0,0,SUMIF({dp_rng},">="&{{c}}{{checkdp}},{n_rng})/{g["total"]})'),
        ("hard", f"=IF({{c}}{{checkdp}}<=0,1,IF({{c}}{{checkdp}}<=6000,{g['pc']},0))"),
        ("die", "=MIN(1,{c}{hard}+IF({c}{barrier}=1,0,{c}{battleloss}))"),
        ("live", "=1-{c}{die}"),
        ("pay", "=IF({c}{barrier}=1,{c}{battleloss},0)"),
    ]
    for t in (0, 1):
        x = "{c}{dp}" if t == 0 else "({c}{dp}-IF({c}{immune}=1,0,{minus}))"
        helpers += [
            (f"x{t}", "=" + x),
            (f"dead{t}", f"=IF(AND({{c}}{{x{t}}}<={g['y']},{{c}}{{barrier}}=0),1,0)"),
            (f"pierce{t}", f"=IF(AND({{c}}{{piercing}}=1,{{c}}{{x{t}}}>={g['y']},{{c}}{{dead{t}}}=0),1,0)"),
            (f"idle{t}", f"=1-{{c}}{{dead{t}}}-{{c}}{{pierce{t}}}"),
            (f"pay{t}", f"=IF(AND({{c}}{{x{t}}}<={g['y']},{{c}}{{barrier}}=1),1,0)"),
        ]
    for key, _ in helpers:
        r += 1
        ENGINE_ATT_ROWS[key] = r
    for key, f in helpers:
        rr = ENGINE_ATT_ROWS[key]
        put(ws, rr, 1, key, F_BOLD)
        for a in range(gc.MAX_ATTACKERS):
            col = get_column_letter(3 + a)
            fmt = {k: f"${v}" for k, v in ENGINE_ATT_ROWS.items()}
            put(ws, rr, 3 + a, f.format(c=col, row=ar + a, minus=g["minus"], **fmt))

    def token(tok):
        if tok[0] == "blk":
            _, field, a, s = tok
            col = get_column_letter(3 + a)
            return f"IF({s}<{g['s0']},{col}${ENGINE_ATT_ROWS[field + '1']},{col}${ENGINE_ATT_ROWS[field + '0']})"
        return f"{get_column_letter(3 + tok[1])}${ENGINE_ATT_ROWS[tok[0]]}"

    def term(src, factors):
        return "*".join([src] + [token(t) for t in factors])

    states = gc.states()
    steps = gc.steps()
    first_col = 6  # F = initial distribution
    for policy in gc.POLICIES:
        r += 3
        put(ws, r, 1, f"Policy: they {'block early' if policy == 'early' else 'save blocks for lethal'}", F_BOLD, fill=FILL_SUB)
        head = r + 1
        act = r + 2
        header(ws, head, 1, ["s", "attacker", "b", "", "", "initial"] + [
            f"A{st[1] + 1} atk{st[2]} " + ("start" if st[0] == "start" else f"chk{st[3]}") for st in steps])
        put(ws, act, 1, "active", F_BOLD)
        row_of = {}
        for i, stt in enumerate(states):
            row_of[stt] = act + 1 + i
        acc_rows = {k: act + 1 + len(states) + i for i, k in enumerate(("checks", "deaths", "pays"))}
        for stt, rr in row_of.items():
            if stt == gc.WON:
                put(ws, rr, 1, "WON", F_BOLD)
            else:
                put(ws, rr, 1, stt[0])
                put(ws, rr, 2, ("dead", "idle", "checking")[stt[1]])
                put(ws, rr, 3, stt[2])
            init = "0" if stt == gc.WON else f"IF(AND({stt[0]}={g['s0']},{stt[1]}={gc.IDLE},{stt[2]}={g['b0']}),1,0)"
            put(ws, rr, first_col, "=" + init)
        for k, rr in acc_rows.items():
            put(ws, rr, 1, k, F_BOLD)
        for ci, stp in enumerate(steps):
            col = first_col + 1 + ci
            L, P = get_column_letter(col), get_column_letter(col - 1)
            a = stp[1]
            ac = get_column_letter(3 + a)
            if stp[0] == "start":
                put(ws, act, col, f"=IF({stp[2]}<={ac}${ENGINE_ATT_ROWS['attacks']},1,0)")
            else:
                put(ws, act, col, f"=IF(AND({stp[2]}<={ac}${ENGINE_ATT_ROWS['attacks']},{stp[3]}<={ac}${ENGINE_ATT_ROWS['checks']}),1,0)")
            incoming = {stt: [] for stt in states}
            accs = {k: [] for k in acc_rows}
            for src in states:
                targets, contrib = gc.transitions(stp, src, policy)
                for tgt, factors in targets:
                    incoming[tgt].append(term(f"{P}{row_of[src]}", factors))
                for kind, factors in contrib.items():
                    accs[kind].append(term(f"{P}{row_of[src]}", factors))
            for stt, terms in incoming.items():
                body = "+".join(terms) if terms else "0"
                put(ws, row_of[stt], col, f"=IF({L}${act}=1,{body},{P}{row_of[stt]})")
            for kind, terms in accs.items():
                body = "+".join(terms) if terms else "0"
                put(ws, acc_rows[kind], col, f"=IF({L}${act}=1,{body},0)")
        last_col = get_column_letter(first_col + len(steps))
        lo, hi = row_of[states[0]], row_of[states[-2]]  # non-WON rows
        sums = {k: f"SUM({get_column_letter(first_col + 1)}{rr}:{last_col}{rr})" for k, rr in acc_rows.items()}
        untouched = f"SUMIF($A${lo}:$A${hi},{g['s0']},{last_col}{lo}:{last_col}{hi})"
        p_any = f"IF({g['s0']}>0,1-{untouched},0)"
        zero_hi = row_of[(0, gc.CHECKING, gc.B_MAX)]
        outs = {
            "p_win": f"={last_col}{row_of[gc.WON]}",
            "p_zero": f"=SUM({last_col}{lo}:{last_col}{zero_hi})",
            "security_left": f"=SUMPRODUCT($A${lo}:$A${hi},{last_col}{lo}:{last_col}{hi})",
            "checks": f"={sums['checks']}",
            "free_tamers": f"={sums['checks']}*{g['pt']}",
            "free_plays": f"={sums['checks']}*{g['po']}",
            "attackers_lost": f"={sums['deaths']}",
            "barrier_paid": f"={sums['pays']}",
            "your_security_after": f"={calc}{CALC_IN['your_security']}-{sums['pays']}-{calc}{CALC_IN['first_check_trash']}*{p_any}",
        }
        r = acc_rows["pays"] + 2
        header(ws, r, 1, ["Output", "Value"], fill=FILL_SUB, font=F_BOLD)
        for field, f in outs.items():
            r += 1
            put(ws, r, 1, field, F_BOLD)
            put(ws, r, 2, f)
            ENGINE_OUT[(policy, field)] = f"$B${r}"
    widths(ws, {"A": 16, "B": 12, "C": 12, "D": 12, "E": 12})
    ws.sheet_state = "hidden"


PROFILE_COLS = ["Value", "Evidence"]


def sheet_profiles(wb, links):
    ws = wb.create_sheet("Opponent Profiles")
    title(ws, "Opponent profiles",
          "What each deck does on its own turns (T2 = their 2nd turn). Hover an evidence cell for its source. 'Was' = the earlier guess.")
    header(ws, 4, 1, ["Parameter"] + [m for m in MATCHUPS for _ in PROFILE_COLS] + ["Was (pre-research, J / H / T / D)", "Basis"])
    for j in range(len(MATCHUPS)):
        ws.merge_cells(start_row=4, start_column=2 + 2 * j, end_row=4, end_column=3 + 2 * j)
    header(ws, 5, 1, [""] + PROFILE_COLS * len(MATCHUPS) + ["", ""], fill=FILL_SUB, font=F_BOLD)
    was_col, basis_col = 2 + 2 * len(MATCHUPS), 3 + 2 * len(MATCHUPS)
    r = 6

    def section(label, text):
        nonlocal r
        put(ws, r, 1, label, F_BOLD, fill=FILL_SUB)
        put(ws, r, 2, text, F_NOTE, fill=FILL_SUB)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=basis_col)
        r += 1

    def was_text(was):
        if was is None:
            return "new row"
        return " / ".join(f"{v:.0%}" if isinstance(v, float) else str(v) for v in was)

    section("Their security (per check we make)", "Live links to the Matchups sheet, counted from the regional lists.")
    for label, key, was in SECURITY_LINKS:
        put(ws, r, 1, label, F_BOLD, wrap=True)
        for j in range(len(MATCHUPS)):
            put(ws, r, 2 + 2 * j, f"=Matchups!{get_column_letter(2 + j)}{links['prob'][key]}", F_LINK, fmt=PCT, align="left")
            put_tag(ws, r, 3 + 2 * j, "COUNTED", "regional top-cut list, Meta Decklists sheet")
        put(ws, r, was_col, was_text(was), F_NOTE)
        r += 1
    for name, text, rows in PROFILE_SECTIONS:
        section(name, text)
        for row in rows:
            put(ws, r, 1, row.label, F_BOLD, wrap=True)
            for j, cell in enumerate(row.cells):
                put(ws, r, 2 + 2 * j, cell.value, F_INPUT, wrap=True, align="left")
                put_tag(ws, r, 3 + 2 * j, cell.tag, cell.source)
            put(ws, r, was_col, was_text(row.was), F_NOTE, wrap=True)
            put(ws, r, basis_col, row.basis, F_NOTE, wrap=True)
            r += 1
    r += 1
    put(ws, r, 1, "How these are used", F_BOLD)
    put(ws, r, 2, "Each opponent turn: roll removal against our board (immunity, Barrier and Habakirimon's protection apply by effect type), roll their attacks into our real security cards, roll blockers against our attacks, and resolve our checks against their security composition. Output: win rate per profile and which play rules change.", wrap=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=basis_col)
    widths(ws, {"A": 40, **{get_column_letter(c): (26 if c % 2 == 0 else 13) for c in range(2, was_col)},
                get_column_letter(was_col): 22, get_column_letter(basis_col): 30})


def sheet_rules(wb):
    ws = wb.create_sheet("Play Rules")
    title(ws, "Play rules")
    header(ws, 4, 1, ["Rule", "What to do", "Number or source behind it", "Evidence"])
    r = 5
    for a, b, c, tag in RULES:
        put(ws, r, 1, a, F_BOLD, wrap=True)
        put(ws, r, 2, b, wrap=True)
        put(ws, r, 3, c, wrap=True)
        put_tag(ws, r, 4, tag)
        r += 1
    r += 1
    header(ws, r, 1, ["Flaw", "Number", "Fix candidate", "Evidence"])
    for a, b, c in FLAWS:
        r += 1
        put(ws, r, 1, a, F_BOLD, wrap=True)
        put(ws, r, 2, b, wrap=True)
        put(ws, r, 3, c, wrap=True)
        put_tag(ws, r, 4, "SIM")
    widths(ws, {"A": 34, "B": 90, "C": 66, "D": 18})


def sheet_playbooks(wb, links):
    for sheet, pb in PLAYBOOKS.items():
        ws = wb.create_sheet(sheet)
        name = MATCHUPS[pb["matchup"]]
        title(ws, f"Glowing Dawn vs {name}", "From real games, pilot guides and card text. Videos are listed on the Sources sheet.")
        wr = links["winrate"][WINRATES[pb["matchup"]][0]]
        put(ws, 4, 1, "Real result (DigiLab)", F_BOLD)
        put(ws, 4, 2, f"=Matchups!E{wr}", F_LINK, fmt=PCT, fill=FILL_KEY)
        put(ws, 4, 3, f'="95% interval "&TEXT(Matchups!F{wr},"0%")&" to "&TEXT(Matchups!G{wr},"0%")&", "&(Matchups!B{wr}+Matchups!C{wr})&" decided games"', F_LINK)
        r = 5
        for label, text in pb["facts"]:
            put(ws, r, 1, label, F_BOLD)
            put(ws, r, 2, text)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            r += 1
        for sec, points in pb["sections"]:
            r += 1
            header(ws, r, 1, [sec, "Point", "Evidence", "Source"])
            for i, (point, tag, src) in enumerate(points, 1):
                r += 1
                put(ws, r, 1, i, align="center")
                put(ws, r, 2, point, wrap=True)
                put_tag(ws, r, 3, tag, src)
                put(ws, r, 4, src, F_NOTE, wrap=True)
        widths(ws, {"A": 26, "B": 105, "C": 16, "D": 40})


SOURCES = [
    ("Card text (all 22 of your cards + 78 meta cards)", "https://digimoncardgame.fandom.com/ (MediaWiki API, pages titled by card number)", "community transcription; verify surprises against the physical card"),
    ("DUAL cards (4-6) and Arts Digivolve (4-20)", "https://world.digimoncard.com/rule/pdf/general_rule.pdf", "Comprehensive Rules dated 2026-09-18"),
    ("EX13 meta shares and win rates", "https://digimon.digilab.cards/meta", ""),
    ("Regionals @ NoHeroes, Dusseldorf, 181 players, 2026-10-04", "https://digimon.digilab.cards/tournament/10314", "top 32 archetypes and lists"),
    ("Regionals @ CCG Eventos, Sao Paulo, 118 players, 2026-09-19", "https://digilab.cards/tournament/9745", "Toho Braves 1st at 65%; Glowing Dawn 15 pilots at 36%"),
    ("Jupitermon 1st (Fish)", META_LINKS["Jupitermon"], ""),
    ("TS Mervamon 5th (UnkindledOne)", META_LINKS["TS Mervamon"], "4x Homeros"),
    ("Toho Braves 10th (Vaelthas)", META_LINKS["Toho Braves"], ""),
    ("DATA SQUAD Ravemon 16th (Rayquon)", META_LINKS["DATA SQUAD"], ""),
    ("Glowing Dawn 4th (Quang-Minh)", META_LINKS["Glowing Dawn"], ""),
    ("Gen Con 2026 Regional (153 players, BT25): Glowing Dawn 3rd/5th/7th/8th, Jupitermon 2nd", "https://digimoncard.io/category/best-digimon-decks/2026", "via search summary; site blocks direct fetch"),
    ("Goldfish simulator", "~/digimon-goldfish (python3 -m goldfish)", "engine, effects, policy; tests under tests/"),
]


def sheet_sources(wb):
    ws = wb.create_sheet("Sources")
    title(ws, "Sources")
    header(ws, 4, 1, ["What", "Where", "Note"])
    r = 4
    for a, b, c in SOURCES + RESEARCH_SOURCES:
        r += 1
        put(ws, r, 1, a, wrap=True)
        put(ws, r, 2, b)
        note(ws, r, 3, c)
    r += 2
    header(ws, r, 1, ["Video (ID as cited in the research)", "Link", ""])
    for vid, what in VIDEOS:
        r += 1
        put(ws, r, 1, f"{vid}: {what}", wrap=True)
        put(ws, r, 2, f"https://www.youtube.com/watch?v={vid}")
    r += 1
    note(ws, r, 1, "Unread high-value videos (YouTube rate-limited the research) are listed per deck at the end of each section in research/round1-timings.md and research/round2-playbooks.md.")
    widths(ws, {"A": 60, "B": 70, "C": 55})


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--out", type=Path, default=ROOT / "reports" / "glowing-dawn-theorycraft.xlsx")
    ap.add_argument("--deck", type=Path, default=ROOT / "decks" / "glowing_dawn.txt")
    ap.add_argument("--games", type=int, default=20000)
    ap.add_argument("--seed", type=int, default=1)
    args = ap.parse_args(argv)
    deck = parse_decklist(args.deck.read_text())
    random.seed(args.seed)
    print("running simulations...", flush=True)
    sims = run_sims(deck, args.games, args.seed)
    stamp = dt.date.today().isoformat()
    wb = Workbook()
    sheet_readme(wb, args.games, args.seed, stamp)
    totals = sheet_decklist(wb, deck)
    sheet_consistency(wb, totals)
    sheet_goldfish(wb, sims, args.seed)
    sheet_board(wb, sims)
    sheet_setup_math(wb)
    links = sheet_matchups(wb)
    sheet_meta_lists(wb)
    sheet_comparison(wb, deck)
    sheet_profiles(wb, links)
    sheet_rules(wb)
    sheet_playbooks(wb, links)
    sheet_calc_engine(wb, links)
    sheet_calculator(wb, links)
    sheet_sources(wb)
    wb.move_sheet("Calculator", offset=1 - wb.sheetnames.index("Calculator"))
    wb.move_sheet(ENGINE, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(ENGINE))
    finalize(wb)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
