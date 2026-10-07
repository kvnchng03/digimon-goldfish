"""Shared spreadsheet look for the workbooks: plain Excel style, autofit rows, header alignment.

Used by tools/build_workbook.py (the Glowing Dawn theorycraft workbook) and
tools/build_player_workbook.py (general Digimon TCG player math).
"""

from __future__ import annotations

import math

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

# ---------------------------------------------------------------------- styling
# Plain Excel look: Calibri, gridlines on, tables boxed in thin gray borders with a darker
# rule under light gray headers, no frozen panes. Blue text = an input you can change.
FONT = "Calibri"
F_BASE = Font(name=FONT, size=11)
F_BOLD = Font(name=FONT, size=11, bold=True)
F_TITLE = Font(name=FONT, size=15, bold=True)
F_HEAD = Font(name=FONT, size=11, bold=True)
F_INPUT = Font(name=FONT, size=11, color="0000FF")
F_LINK = Font(name=FONT, size=11, color="375623")
F_NOTE = Font(name=FONT, size=10, italic=True, color="595959")
F_TAG = Font(name=FONT, size=9)
FILL_HEAD = PatternFill("solid", fgColor="F2F2F2")
FILL_SUB = PatternFill("solid", fgColor="F2F2F2")
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=Side(style="thin", color="595959"))
PCT = "0.0%"
NUM1 = "0.0"
NUM2 = "0.00"


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_NOTE


def header(ws, row, col, values, fill=FILL_HEAD, font=F_HEAD):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = font
        c.fill = fill
        c.border = HEAD_BORDER
        c.alignment = Alignment(wrap_text=True, vertical="bottom")


def put(ws, row, col, value, font=F_BASE, fmt=None, fill=None, wrap=False, comment=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    c.border = BORDER
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    if comment:
        c.comment = Comment(comment, "goldfish")
    return c


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def text_lines(text: str, width: float, font_size: float = 11) -> int:
    """Lines a wrapped text needs in a cell `width` characters wide (Excel column units)."""
    chars = width * 1.2 * 10 / font_size
    return sum(max(1, math.ceil(len(part) / chars)) for part in text.split("\n"))


def row_height(lines: int, line_pt: int = 13) -> float:
    return line_pt * lines + 3


def autofit_rows(ws, first=1, last=None, line_pt=13):
    """Size rows so wrapped text shows in full. Estimates lines from text length and column
    width (merged cells use their combined width); never shrinks a row below one line.
    Formula cells are skipped: size those yourself with text_lines / row_height."""
    width_of = lambda col: ws.column_dimensions[get_column_letter(col)].width or 8.43  # noqa: E731
    span, covered = {}, set()
    for rng in ws.merged_cells.ranges:
        span[(rng.min_row, rng.min_col)] = sum(width_of(c) for c in range(rng.min_col, rng.max_col + 1))
        covered |= {(r, c) for r in range(rng.min_row, rng.max_row + 1) for c in range(rng.min_col, rng.max_col + 1)}
    for row in ws.iter_rows(min_row=first, max_row=last or ws.max_row):
        lines = 1
        for cell in row:
            key = (cell.row, cell.column)
            if cell.value is None or not cell.alignment.wrap_text or (key in covered and key not in span):
                continue
            text = str(cell.value)
            if text.startswith("="):
                continue
            lines = max(lines, text_lines(text, span.get(key, width_of(cell.column)), cell.font.sz or 10))
        if lines > 1:
            ws.row_dimensions[row[0].row].height = row_height(lines, line_pt)


def align_headers(ws):
    """Right-align a header cell when the column under it holds numbers, so labels sit over values."""
    for row in ws.iter_rows():
        for c in row:
            if c.border.bottom is None or c.border.bottom.color is None or c.border.bottom.color.rgb != HEAD_BORDER.bottom.color.rgb:
                continue
            below = ws.cell(row=c.row + 1, column=c.column)
            v = below.value
            numeric = isinstance(v, int | float) and not isinstance(v, bool)
            formula = isinstance(v, str) and v.startswith("=") and below.alignment.horizontal not in ("left", "center")
            if (numeric and below.alignment.horizontal not in ("left", "center")) or formula:
                c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="right")


def note(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_NOTE
    c.alignment = Alignment(wrap_text=False)
    return c


def link(ws, row, col, sheet, text=None):
    """A clickable cell that jumps to another sheet."""
    c = put(ws, row, col, text or sheet, Font(name=FONT, size=11, color="0563C1", underline="single"))
    c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{sheet}'!A1", display=text or sheet)
    return c


def finalize(wb):
    """Last pass over every sheet: size rows to their text, line headers up with numbers,
    and make Excel recalculate on open (openpyxl stores formulas without values)."""
    for ws in wb.worksheets:
        autofit_rows(ws)
        align_headers(ws)
    wb.calculation.fullCalcOnLoad = True
