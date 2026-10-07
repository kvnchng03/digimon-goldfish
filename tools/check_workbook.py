"""Recalculate every formula in the workbook and report errors.

    python3 tools/check_workbook.py [reports/glowing-dawn-theorycraft.xlsx] [--show Sheet!A1 ...]
    python3 tools/check_workbook.py --render DIR [--sheets "Opponent Profiles" ...] [--recalc] [--scenario NAME]

openpyxl writes formulas without cached values, and Excel automation is unreliable here,
so this evaluates them with the `formulas` library (pip install formulas). Exit code 1 if
any formula evaluates to an Excel error (#REF!, #DIV/0!, #NAME?, ...).

It then cross-checks the file against Python: the Glowing Dawn workbook's Calculator is
driven through several situations and compared with goldfish.calc; the player-math workbook
and the player log are compared cell by cell with the values build_player_workbook expects,
and the builder's scenarios (puzzle picks, edge-case inputs) are typed in and checked too.

--render writes one PNG per sheet via macOS Quick Look (qlmanage) to eyeball layout:
row heights, wrapping, alignment. Quick Look does not recalculate, so formula cells show 0
unless --recalc bakes the evaluated values in; --scenario types in one of the player
workbook's scenarios first (say, every puzzle answered) to see the revealed state.
"""

from __future__ import annotations

import argparse
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import formulas
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))

import build_player_workbook as bpw  # noqa: E402
import build_workbook as bw  # noqa: E402

from goldfish import calc as gc  # noqa: E402

# Situations the cross-check sets on the Calculator: inputs, then attacker rows
# (name, DP, checks, attacks, Barrier, immune, Piercing).
SCENARIOS = [
    ({}, None),  # the defaults the builder wrote: Jupitermon at 2, Wrath Mode blocking, Holy out
    ({"matchup": bw.MATCHUPS[2], "their_security": 3, "stoppers": 1, "stopper_dp": 12000, "minus": 0, "first_check_trash": 0},
     [("Atratusmon", 12000, 2, 2, "N", "Y", "N"), ("Habakirimon", 12000, 1, 2, "N", "N", "N"), ("Armalizamon", 4000, 1, 1, "N", "N", "Y")]),
    ({"matchup": bw.MATCHUPS[3], "their_security": 5, "stoppers": 0, "minus": 0, "first_check_trash": 0},
     [("A", 15000, 2, 2, "N", "Y", "N"), ("B", 7000, 1, 1, "Y", "N", "N"), ("C", 6000, 3, 1, "N", "N", "N")]),
    ({"matchup": bw.MATCHUPS[1], "their_security": 1, "stoppers": 1, "stopper_dp": 15000, "minus": 2000},
     [("A", 16000, 1, 2, "N", "N", "Y"), ("B", 12000, 2, 1, "Y", "Y", "Y"), ("C", 0, 1, 0, "N", "N", "N")]),
]


def _combin(n, k):
    return math.comb(int(n), int(k))


def register_missing():
    """Functions the workbook uses that `formulas` lacks."""
    fns = formulas.get_functions()
    if "COMBIN" not in fns:
        fns["COMBIN"] = formulas.functions.wrap_ufunc(np.vectorize(_combin, otypes=[float]))


def scalar(v):
    while hasattr(v, "value") and not isinstance(v, str | int | float):
        v = v.value
    if isinstance(v, np.ndarray):
        v = v.ravel()[0]
    return v


def typed_copy(path: Path, typed: dict, tmp: str, stem: str) -> Path:
    """A copy of the workbook with `typed` ({"Sheet!CELL": value}) typed in."""
    wb = load_workbook(path)
    for ref, v in typed.items():
        sheet, cell = ref.rsplit("!", 1)
        wb[sheet][cell] = v
    copy = Path(tmp) / f"{stem}.xlsx"
    wb.save(copy)
    return copy


def baked_copy(path: Path, tmp: str) -> Path:
    """A copy whose formulas are replaced by their evaluated values, so Quick Look shows them."""
    register_missing()
    values = evaluate(path)
    wb = load_workbook(path)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = values.get(f"{ws.title}!{c.coordinate}".upper())
                    c.value = v if isinstance(v, str | int | float) and not isinstance(v, bool) else None
    copy = Path(tmp) / "baked.xlsx"
    wb.save(copy)
    return copy


def render(path: Path, out: Path, sheets: list[str], recalc: bool = False, scenario: str | None = None) -> None:
    """Quick Look thumbnails only show the active first sheet, so render a copy per sheet.
    `recalc` bakes evaluated values in first (Quick Look doesn't calculate); `scenario` types in
    a player-workbook scenario from the builder (matched by name) before that."""
    out.mkdir(parents=True, exist_ok=True)
    names = sheets or [ws.title for ws in load_workbook(path).worksheets if ws.sheet_state == "visible"]
    with tempfile.TemporaryDirectory() as tmp:
        if scenario:
            bpw.SCENARIOS.clear()
            bpw.build(Path(tmp) / "expected.xlsx")
            name, typed, _ = next(sc for sc in bpw.SCENARIOS if scenario in sc[0])
            print(f"scenario: {name}")
            path = typed_copy(path, typed, tmp, "typed")
        if recalc or scenario:
            path = baked_copy(path, tmp)
        for name in names:
            wb = load_workbook(path)
            ws = wb[name]
            wb.move_sheet(ws, offset=-wb.index(ws))
            ws.sheet_state = "visible"
            wb.active = 0
            for other in wb.worksheets:
                other.sheet_view.tabSelected = other is ws
            stem = name.replace(" ", "_")
            copy = Path(tmp) / f"{stem}.xlsx"
            wb.save(copy)
            subprocess.run(["qlmanage", "-t", "-s", "1800", "-o", tmp, str(copy)], capture_output=True, check=True)
            shutil.move(Path(tmp) / f"{stem}.xlsx.png", out / f"{stem}.png")
            print(out / f"{stem}.png")


def evaluate(path: Path) -> dict:
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    return {key.split("]", 1)[-1].replace("'", "").upper(): scalar(rng) for key, rng in sol.items()}


def situation_from(ws) -> gc.Situation:
    """The Calculator's inputs as a goldfish.calc Situation (lists as the builder counted them)."""
    val = lambda key: ws[bw.CALC_IN[key]].value  # noqa: E731
    lists = [c for h, c in bw.META_LISTS.items() if not h.startswith("Glowing Dawn")]
    dps, tamers, crimson, other = bw.security_composition(lists[bw.MATCHUPS.index(val("matchup"))])
    attackers = []
    for r in range(bw.CALC_ATTACKER_ROW, bw.CALC_ATTACKER_ROW + gc.MAX_ATTACKERS):
        cell = lambda k: ws[f"{bw.CALC_ATTACKER_COLS[k]}{r}"].value  # noqa: E731, B023
        attackers.append(gc.Attacker(dp=cell("dp"), checks=cell("checks"), attacks=cell("attacks"),
                                     barrier=cell("barrier") == "Y", immune=cell("immune") == "Y",
                                     piercing=cell("piercing") == "Y"))
    return gc.Situation(gc.Security(tuple(sorted(dps.items())), tamers, crimson, other),
                        their_security=val("their_security"), attackers=tuple(attackers), stoppers=val("stoppers"),
                        stopper_dp=val("stopper_dp"), minus=val("minus"), your_security=val("your_security"),
                        first_check_trash=val("first_check_trash"))


def crosscheck(path: Path) -> list[str]:
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        for i, (inputs, attackers) in enumerate(SCENARIOS):
            wb = load_workbook(path)
            ws = wb["Calculator"]
            for key, v in inputs.items():
                ws[bw.CALC_IN[key]] = v
            for r, row in enumerate(attackers or []):
                for col, v in zip(bw.CALC_ATTACKER_COLS.values(), row, strict=True):
                    ws[f"{col}{bw.CALC_ATTACKER_ROW + r}"] = v
            copy = Path(tmp) / f"scenario{i}.xlsx"
            wb.save(copy)
            values = evaluate(copy)
            sit = situation_from(ws)
            for j, policy in enumerate(gc.POLICIES):
                want = gc.solve(sit, policy)
                for k, (label, field, _) in enumerate(bw.CALC_RESULTS):
                    cell = f"CALCULATOR!{'CD'[j]}{bw.CALC_RESULT_ROW + k}"
                    got = values.get(cell)
                    if got is None or abs(float(got) - getattr(want, field)) > 1e-9:
                        problems.append(f"scenario {i} {policy} {label}: sheet {got!r} vs model {getattr(want, field)!r}")
            print(f"scenario {i}: P(win) early {gc.solve(sit, 'early').p_win:.3f} / lethal {gc.solve(sit, 'lethal').p_win:.3f}")
    return problems


def compare(values: dict, want: dict, label: str = "") -> list[str]:
    problems = []
    for cell, expected in want.items():
        got = values.get(cell)
        if isinstance(expected, bool) or not isinstance(expected, int | float):
            same = got == expected
        elif isinstance(got, int | float) and not isinstance(got, bool):
            same = abs(float(got) - expected) < 1e-9
        else:
            same = False
        if not same:
            problems.append(f"{label}{cell}: sheet {got!r} vs Python {expected!r}")
    return problems


def expected_values(values: dict, kind: str) -> list[str]:
    """Player-math workbook or log: compare cells with the values Python computes for the defaults."""
    bpw.EXPECT.clear()
    bpw.SCENARIOS.clear()
    with tempfile.TemporaryDirectory() as tmp:
        if kind == "player":
            bpw.build(Path(tmp) / "expected.xlsx")
        else:
            bpw.build_log(Path(tmp) / "expected-log.xlsx", force=True)
    return compare(values, bpw.EXPECT)


def player_scenarios(path: Path) -> list[str]:
    """Type each builder scenario's inputs into a copy of the workbook, recalculate, compare."""
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        for i, (name, typed, want) in enumerate(bpw.SCENARIOS):
            found = compare(evaluate(typed_copy(path, typed, tmp, f"scenario{i}")), want, f"[{name}] ")
            print(f"scenario '{name}': {len(want)} cells, {len(found)} mismatches")
            problems += found
    return problems


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("path", nargs="?", type=Path, default=ROOT / "reports" / "glowing-dawn-theorycraft.xlsx")
    ap.add_argument("--show", nargs="*", default=[], help="cells to print, e.g. Matchups!E5")
    ap.add_argument("--render", type=Path, help="write a PNG per sheet into this directory and exit")
    ap.add_argument("--sheets", nargs="*", default=[], help="with --render: only these sheets")
    ap.add_argument("--recalc", action="store_true", help="with --render: show evaluated values instead of 0")
    ap.add_argument("--scenario", help="with --render: type in the player-workbook scenario whose name contains this")
    args = ap.parse_args(argv)
    if args.render:
        render(args.path, args.render, args.sheets, args.recalc, args.scenario)
        return 0
    register_missing()

    wb = load_workbook(args.path)
    formula_cells = {f"{ws.title}!{c.coordinate}".upper() for ws in wb for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith("=")}
    values = evaluate(args.path)

    errors = []
    for cell in sorted(formula_cells):
        v = values.get(cell)
        if v is None or isinstance(v, formulas.functions.XlError) or (isinstance(v, str) and v.startswith("#")):
            errors.append((cell, v))
    print(f"{len(formula_cells)} formulas, {len(errors)} errors")
    for cell, v in errors:
        print(f"  {cell}: {v!r}")
    for cell in args.show:
        print(f"{cell} = {values.get(cell.upper())!r}")
    sheets = load_workbook(args.path, read_only=True).sheetnames
    if "Calculator" in sheets:
        problems = crosscheck(args.path)
        print(f"calculator cross-check: {len(SCENARIOS)} situations, {len(problems)} mismatches")
    else:
        kind = "player" if "Quiz" in sheets else "log"
        problems = expected_values(values, kind)
        print(f"expected-value check ({kind}): {len(bpw.EXPECT)} cells, {len(problems)} mismatches")
        if kind == "player":
            problems += player_scenarios(args.path)
    for line in problems:
        print(f"  {line}")
    return 1 if errors or problems else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
